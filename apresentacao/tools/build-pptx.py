#!/usr/bin/env python3
"""
Build a PPTX from the rendered slide PNGs.

Each slide is placed as a full-bleed background image in a 16:10 widescreen
PowerPoint. The user can then edit in PowerPoint/Keynote — most commonly
to overlay the official logos (just Insert > Picture and place on top).
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Emu

SRC_DIR = Path("/workspace/apresentacao/pptx-source")
OUT     = Path("/workspace/apresentacao/casting-blaze-proposta.pptx")

# 16:10 widescreen — matches our 1280x800 slide canvas
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(8.333)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

blank_layout = prs.slide_layouts[6]  # fully blank layout

for i in range(1, 13):
    img = SRC_DIR / f"s{i:02d}.png"
    if not img.exists():
        print(f"missing {img}")
        continue
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(
        str(img),
        left=Emu(0), top=Emu(0),
        width=SLIDE_W, height=SLIDE_H,
    )
    print(f"added slide {i}")

prs.save(OUT)
print(f"\nPPTX written to {OUT}")
print(f"size: {OUT.stat().st_size / 1024 / 1024:.1f} MB")
