#!/usr/bin/env python3
"""
Gera um PPTX EDITAVEL do deck Casting Blaze, com todos os textos como
text boxes nativos do PowerPoint (e nao como imagem). O usuario pode
editar qualquer frase, numero, titulo direto no PowerPoint/Keynote.

Trade-off vs. a versao "imagem":
- Vantagem: tudo editavel
- Desvantagem: alguns efeitos do HTML/CSS (gradientes, blur, glow text)
  ficam simplificados — usamos cores solidas e shapes basicos do PPTX.

Layout: 16:10 widescreen mapeando nosso canvas 1280x800px.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

OUT = Path("/workspace/apresentacao/casting-blaze-proposta-editavel.pptx")
IMG_DIR = Path("/workspace/apresentacao/assets/images")

# ============================================================================
# Helpers de unidade
# ============================================================================
# Mapeamento: 1 px do nosso canvas HTML (1280x800) = 9525 EMU
#   1 inch = 914400 EMU = 96 px
#   slide 1280 px = 13.333 in = 12_192_000 EMU
def px(v): return int(v * 9525)
def pt(v): return Pt(v)
# CSS px → pt aproximado (1pt = 1.333 px)
def fpt(p): return Pt(p * 0.75)


# ============================================================================
# Paleta
# ============================================================================
BG_DARK    = RGBColor(0x0A, 0x01, 0x03)
BG_MID     = RGBColor(0x15, 0x04, 0x0A)
BG_CARD    = RGBColor(0x14, 0x03, 0x09)
RED_BLAZE  = RGBColor(0xE5, 0x0E, 0x2B)
RED_SOFT   = RGBColor(0xFF, 0x3B, 0x54)
RED_DEEP   = RGBColor(0x4A, 0x0A, 0x16)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
WHITE_DIM  = RGBColor(0xC8, 0xC4, 0xC4)
WHITE_MUTE = RGBColor(0x8A, 0x86, 0x86)
WATERMARK  = RGBColor(0x1E, 0x14, 0x16)

FONT       = 'Arial'
FONT_BOLD  = 'Arial'
FONT_BLACK = 'Arial Black'


# ============================================================================
# Setup do deck
# ============================================================================
prs = Presentation()
prs.slide_width  = px(1280)
prs.slide_height = px(800)
BLANK = prs.slide_layouts[6]


# ============================================================================
# Helpers de elementos
# ============================================================================
def set_bg(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color

def remove_shadow(shape):
    """Remove sombra padrao do PPTX que polui visualmente."""
    sp = shape.shadow
    try:
        sp.inherit = False
    except Exception:
        pass
    # tentar tambem via XML
    try:
        sppr = shape._element.spPr
        ext = sppr.find(qn('a:effectLst'))
        if ext is not None:
            sppr.remove(ext)
        # adiciona effectLst vazio (sem efeito)
        from pptx.oxml.ns import qn as _qn
        new = etree.SubElement(sppr, _qn('a:effectLst'))
    except Exception:
        pass

from pptx.oxml.ns import qn

def add_text(slide, x, y, w, h, text, *,
             size=14, color=WHITE, bold=False, italic=False,
             font=FONT, align='left', vanchor='top', spacing=0,
             line_spacing=1.2):
    """Adiciona um text box editavel. Aceita lista de tuplas para
    multiplos runs com formatacoes diferentes:
        text = [
          ("Olá ", {}),
          ("mundo", {"bold": True, "color": RED_BLAZE}),
        ]
    """
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    if vanchor == 'middle':
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif vanchor == 'bottom':
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM

    def apply_run(run, t, st):
        run.text = t
        run.font.name = st.get('font', font)
        run.font.size = Pt(st.get('size', size))
        run.font.color.rgb = st.get('color', color)
        run.font.bold = st.get('bold', bold)
        run.font.italic = st.get('italic', italic)
        sp = st.get('spacing', spacing)
        if sp:
            rPr = run._r.get_or_add_rPr()
            rPr.set('spc', str(int(sp * 100)))

    p = tf.paragraphs[0]
    if align == 'center':
        p.alignment = PP_ALIGN.CENTER
    elif align == 'right':
        p.alignment = PP_ALIGN.RIGHT
    p.line_spacing = line_spacing

    if isinstance(text, str):
        r = p.add_run()
        apply_run(r, text, {})
    elif isinstance(text, list):
        # tuple = (text, style); pode ter "\n" pra quebra explicita
        first = True
        for item in text:
            t, st = item if isinstance(item, tuple) else (item, {})
            if t == '\n':
                p = tf.add_paragraph()
                if align == 'center':
                    p.alignment = PP_ALIGN.CENTER
                elif align == 'right':
                    p.alignment = PP_ALIGN.RIGHT
                p.line_spacing = line_spacing
                first = True
                continue
            r = p.add_run()
            apply_run(r, t, st)
            first = False
    return tb

def add_rect(slide, x, y, w, h, *, fill=None, line=None, line_w=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.text_frame.text = ''
    remove_shadow(s)
    return s

def add_card(slide, x, y, w, h, *,
             fill=BG_CARD, border=RED_BLAZE, border_w=0.75, radius=0.06):
    """Card retangular arredondado com borda fina vermelha."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                px(x), px(y), px(w), px(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(border_w)
    s.adjustments[0] = radius
    s.text_frame.text = ''
    remove_shadow(s)
    return s

def add_pill(slide, x, y, w, h, text, *,
             fill=None, border=RED_BLAZE, text_color=None):
    """Pill (rounded rect totalmente arredondado)."""
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(1)
    s.adjustments[0] = 0.5
    tf = s.text_frame
    tf.margin_left = tf.margin_right = px(8)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = text_color or (WHITE if fill else RED_BLAZE)
    rPr = r._r.get_or_add_rPr()
    rPr.set('spc', '200')
    remove_shadow(s)
    return s

def add_circle(slide, x, y, d, *, fill=None, border=RED_BLAZE,
               text='', text_size=10, text_color=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(d), px(d))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = border
    s.line.width = Pt(0.75)
    tf = s.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT_BLACK
    r.font.size = Pt(text_size)
    r.font.bold = True
    r.font.color.rgb = text_color
    remove_shadow(s)
    return s


# ============================================================================
# Watermark e chrome base de cada slide
# ============================================================================
def add_watermark(slide):
    tb = slide.shapes.add_textbox(px(-40), px(60), px(900), px(560))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, word in enumerate(['Spin', 'GAMING']):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.line_spacing = 0.78
        r = p.add_run()
        r.text = word
        r.font.name = FONT_BLACK
        r.font.size = Pt(180)
        r.font.color.rgb = WATERMARK

def add_chrome(slide, page, *, tagline='', caption=''):
    # top hairline
    add_rect(slide, 0, 14, 1280, 1, fill=RED_BLAZE)
    # bottom hairline
    add_rect(slide, 0, 786, 1280, 1, fill=RED_BLAZE)
    # date
    add_text(slide, 64, 26, 220, 18, 'JUNHO 2026',
             size=8, color=WHITE, bold=True, spacing=3)
    # tagline center
    if tagline:
        add_text(slide, 380, 22, 520, 28, tagline,
                 size=8, color=WHITE, bold=True, align='center',
                 spacing=2, line_spacing=1.4)
    # Spin Gaming logo placeholder (top right, editavel)
    add_text(slide, 1130, 22, 90, 28,
             [('spin', {'size': 16, 'bold': True, 'italic': True, 'font': FONT_BLACK}),
              ('\n', {}),
              ('GAMING', {'size': 6, 'bold': True, 'spacing': 4})],
             color=WHITE, align='right', line_spacing=1.0)
    # footer caption
    add_text(slide, 64, 770, 700, 14, caption.upper(),
             size=7, color=WHITE_MUTE, bold=True, spacing=2)
    # page number
    add_text(slide, 1080, 770, 140, 14, f'{page:02d} / 12',
             size=7, color=WHITE_MUTE, bold=True, align='right', spacing=2)

def add_section_tabs(slide, *, casting_active=False, cashback_active=False):
    """As 2 tabs no topo (CASTING BLAZE / CASHBACK 10%)."""
    x = 410
    y = 60
    for label, active in [('CASTING BLAZE', casting_active),
                          ('CASHBACK 10%', cashback_active)]:
        fill = RGBColor(0x33, 0x05, 0x0E) if active else None
        text_color = WHITE if active else WHITE_MUTE
        add_pill(slide, x, y, 200, 26, label,
                 fill=fill if active else None,
                 border=RED_BLAZE if active else WHITE_MUTE,
                 text_color=text_color)
        x += 215

def new_slide(page, *, tagline='', caption=''):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, BG_MID)
    add_watermark(s)
    add_chrome(s, page, tagline=tagline, caption=caption)
    return s


def add_eyebrow(slide, x, y, w, text):
    """Eyebrow vermelho com hairline embaixo."""
    add_text(slide, x, y, w, 16, text, size=9, color=RED_BLAZE,
             bold=True, spacing=3)
    add_rect(slide, x, y + 22, w, 1, fill=RGBColor(0x7A, 0x0A, 0x18))

def add_bullet_list(slide, x, y, w, items, *, item_h=20, size=10):
    """Lista de bullets com quadradinho vermelho."""
    for i, txt in enumerate(items):
        yy = y + i * item_h
        add_rect(slide, x, yy + 6, 7, 7, fill=RED_BLAZE)
        add_text(slide, x + 16, yy, w - 16, item_h, txt,
                 size=size, color=WHITE_DIM)


# ============================================================================
# SLIDES
# ============================================================================

# ---------------------------------------------------------------------- 01
def slide_01():
    s = new_slide(1,
        tagline='UMA PROPOSTA QUE TRANSFORMA\nINFLUÊNCIA EM VOLUME DE JOGO',
        caption='SPIN GAMING  ×  BLAZE')
    # Blaze wordmark gigante (placeholder editavel - substituir por logo oficial)
    add_text(s, 0, 280, 1280, 200, 'blaze',
             size=140, color=RED_BLAZE, bold=True, align='center',
             font=FONT_BLACK, line_spacing=1.0)
    # Casting Blaze title
    add_text(s, 0, 510, 1280, 40, 'CASTING BLAZE',
             size=24, color=WHITE, bold=True, align='center', spacing=2)
    # Subtitle
    add_text(s, 0, 555, 1280, 30,
             'INFLUÊNCIA  ·  LIVE CASINO  ·  RETENÇÃO',
             size=10, color=WHITE, bold=True, align='center', spacing=6)


# ---------------------------------------------------------------------- 02
def slide_02():
    s = new_slide(2, tagline='CONTEXTO DA PROPOSTA',
                  caption='SPIN GAMING  ×  BLAZE  ·  CONTEXTO')
    # eyebrow
    add_eyebrow(s, 64, 100, 200, 'POR QUE AGORA')
    # headline
    add_text(s, 64, 140, 580, 270,
             [('A SPIN JÁ ENTREGA\nAS ', {}),
              ('MESAS AO VIVO', {'color': RED_BLAZE}),
              ('\nDA BLAZE. AGORA\nQUEREMOS LEVAR\nTRÁFEGO ATÉ ELAS.', {})],
             size=36, color=WHITE, bold=True, font=FONT_BLACK,
             line_spacing=1.05)
    # body 1
    add_text(s, 64, 440, 580, 90,
             ('Hoje a Blaze opera com mesas Spin Gaming integradas no lobby. '
              'A operação é estável, a tecnologia é nossa, o estúdio é nosso '
              'e o treinamento de dealers é feito pela Spin Academy. O próximo '
              'passo natural é ativar a base: trazer mais UAP para essas '
              'mesas e aumentar a frequência de depósitos.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)
    # body 2
    add_text(s, 64, 555, 580, 50,
             'Esta proposta combina duas alavancas que conversam diretamente '
             'com o público brasileiro de cassino ao vivo:',
             size=10.5, color=WHITE_DIM, line_spacing=1.5)
    # pills
    add_pill(s, 64, 615, 200, 28, '01 · CASTING BLAZE',
             fill=RED_BLAZE, text_color=WHITE)
    add_pill(s, 274, 615, 230, 28, '02 · CASHBACK 10% MESAS SPIN')
    # photo
    img = IMG_DIR / 'mesa-blaze-blackjack.jpg'
    if img.exists():
        s.shapes.add_picture(str(img), px(700), px(140), px(516), px(495))


# ---------------------------------------------------------------------- 03
def slide_03():
    s = new_slide(3, tagline='DOIS PILARES, UMA META: VOLUME',
                  caption='SPIN GAMING  ×  BLAZE  ·  PILARES')
    add_section_tabs(s, casting_active=True, cashback_active=True)

    # Card 1 — Casting Blaze
    add_card(s, 64, 120, 565, 620)
    add_text(s, 88, 145, 200, 16, 'PILAR 01',
             size=9, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 88, 165, 500, 32, 'CASTING BLAZE',
             size=20, color=WHITE, bold=True, font=FONT_BLACK)
    add_text(s, 88, 215, 510, 130,
             [('Treinamos influenciadores reconhecidos pelo público da Blaze na '
               'Spin Academy e os colocamos para atuar como dealers oficiais '
               'em mesas Spin × Blaze. Eles transmitem em seus próprios canais '
               'e direcionam audiência diretamente para o lobby do live casino.',
               {})],
             size=10.5, color=WHITE_DIM, line_spacing=1.5)
    add_text(s, 88, 565, 200, 16, 'O QUE ENTREGA',
             size=8, color=WHITE_MUTE, bold=True, spacing=2)
    add_bullet_list(s, 88, 590, 510, [
        'Aquisição de UAP direto do público do influenciador',
        'Brand love via co-criação com nomes que a base já segue',
        'Conteúdo viral com cortes e clipes orgânicos',
        'Tráfego concentrado em mesa específica → GGR mensurável',
    ], item_h=22, size=10)

    # Card 2 — Cashback
    add_card(s, 651, 120, 565, 620)
    add_text(s, 675, 145, 200, 16, 'PILAR 02',
             size=9, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 675, 165, 530, 32, 'CASHBACK 10% MESAS SPIN',
             size=20, color=WHITE, bold=True, font=FONT_BLACK)
    add_text(s, 675, 215, 510, 130,
             ('Programa de devolução semanal de 10% das perdas líquidas '
              'apenas para apostas realizadas nas mesas Spin Gaming dentro '
              'da Blaze. Mecânica simples, transparente, sem rollover abusivo '
              'e em conformidade com a regulação brasileira de iGaming.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)
    add_text(s, 675, 565, 200, 16, 'O QUE ENTREGA',
             size=8, color=WHITE_MUTE, bold=True, spacing=2)
    add_bullet_list(s, 675, 590, 510, [
        'Retenção +20% a +30% (benchmarks globais)',
        'Frequência de depósito +18% (case Vegasino · Optimove)',
        'Ticket médio +25% em jogadores aderentes',
        'Concentração da base ativa nas mesas Spin → mais GGR',
    ], item_h=22, size=10)


# ---------------------------------------------------------------------- 04
def slide_04():
    s = new_slide(4, tagline='PILAR 01 · O CONCEITO',
                  caption='SPIN GAMING  ×  BLAZE  ·  CASTING BLAZE')
    add_section_tabs(s, casting_active=True)

    # photo left
    img = IMG_DIR / 'mesa-blaze-real.png'
    if img.exists():
        s.shapes.add_picture(str(img), px(64), px(120), px(545), px(560))
    # pill on photo
    add_pill(s, 84, 140, 140, 26, 'FORMATO INÉDITO',
             fill=RED_BLAZE, text_color=WHITE)

    # right column
    add_eyebrow(s, 651, 130, 200, 'O CONCEITO')
    add_text(s, 651, 170, 540, 200,
             [('O INFLUENCIADOR QUE O\nJOGADOR JÁ SEGUE,\n', {}),
              ('DENTRO DA MESA', {'color': RED_BLAZE}),
              ('\nDA BLAZE.', {})],
             size=28, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    add_text(s, 651, 400, 540, 100,
             ('Em vez de um banner ou de um link de afiliado, o jogador '
              'encontra o influenciador operando ao vivo a mesa de Blackjack, '
              'Roleta ou Baccarat que está dentro da Blaze. A experiência funde '
              'entretenimento, autoridade do criador e a mesa real onde a aposta '
              'acontece.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)

    add_text(s, 651, 510, 540, 90,
             ('É o único formato do mercado em que o influenciador não recomenda '
              'a casa — ele faz parte da operação dela. Isso elimina o atrito entre '
              'conteúdo e produto e encurta dramaticamente o caminho até o depósito.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)

    # 2 mini cards
    add_card(s, 651, 620, 260, 90)
    add_text(s, 671, 635, 220, 14, 'AUTENTICIDADE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 671, 656, 220, 48,
             'O criador foi formado pela Spin Academy e domina as regras. Não é "encenação".',
             size=9, color=WHITE_DIM, line_spacing=1.4)

    add_card(s, 931, 620, 260, 90)
    add_text(s, 951, 635, 220, 14, 'PROXIMIDADE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 951, 656, 220, 48,
             'Interage com o chat em tempo real, gerando até 4h de permanência por sessão.',
             size=9, color=WHITE_DIM, line_spacing=1.4)


# ---------------------------------------------------------------------- 05
def slide_05():
    s = new_slide(5, tagline='PILAR 01 · COMO FUNCIONA',
                  caption='SPIN GAMING  ×  BLAZE  ·  PROCESSO')
    add_section_tabs(s, casting_active=True)

    # eyebrow + title left
    add_eyebrow(s, 64, 130, 360, 'PROCESSO SPIN ACADEMY → BLAZE LOBBY')
    add_text(s, 64, 170, 620, 130,
             [('4 ETAPAS PARA TRANSFORMAR\nAUDIÊNCIA EM ', {}),
              ('VOLUME REAL', {'color': RED_BLAZE}),
              ('.', {})],
             size=28, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    # 4 steps
    steps = [
        ('1', 'CASTING',
         'Selecionamos junto com a Blaze 3 a 5 influenciadores que já '
         'conversam com o público de cassino — com base em CCV médio, '
         'engajamento real e fit de marca.'),
        ('2', 'SPIN ACADEMY',
         'Treinamento intensivo de 1 a 5 dias (dependendo do jogo): regras '
         'de Blackjack, Roleta e Baccarat, procedimentos operacionais, '
         'dicção e padrão de transmissão Spin.'),
        ('3', 'SHOW AO VIVO',
         'Mesa dedicada com branding Blaze, transmissão simultânea no canal '
         'do influenciador (Kick / YouTube / Instagram) e no lobby da Blaze.'),
        ('4', 'CONTEÚDO & MENSURAÇÃO',
         'Cortes virais, mesa-código exclusiva por criador para tracking de '
         'UAP, FTD e GGR atribuído. Dashboards compartilhados em D+1.'),
    ]
    y = 320
    for num, title, body in steps:
        add_circle(s, 64, y, 38, fill=RGBColor(0x33, 0x05, 0x0E),
                   text=num, text_size=14)
        add_text(s, 116, y, 580, 18, title,
                 size=10, color=RED_BLAZE, bold=True, spacing=1)
        add_text(s, 116, y + 22, 580, 60, body,
                 size=9.5, color=WHITE_DIM, line_spacing=1.45)
        y += 90

    # right sidebar cards
    # jogos sugeridos
    add_card(s, 770, 130, 446, 110)
    add_text(s, 790, 148, 350, 14, 'JOGOS SUGERIDOS',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 790, 172, 200, 22, '• Blackjack VIP',
             size=11, color=WHITE_DIM)
    add_text(s, 990, 172, 200, 22, '• Roleta Brasileira',
             size=11, color=WHITE_DIM)
    add_text(s, 790, 200, 200, 22, '• Baccarat',
             size=11, color=WHITE_DIM)

    # canais de transmissao
    add_card(s, 770, 256, 446, 220)
    add_text(s, 790, 274, 350, 14, 'CANAIS DE TRANSMISSÃO',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    channels = [
        '▶  Kick — média 600k+ live viewers/dia',
        '▶  YouTube Lives + cortes evergreen',
        '▶  Instagram Reels + Stories direcionando para Blaze',
    ]
    for i, ch in enumerate(channels):
        add_text(s, 790, 304 + i * 50, 410, 32, ch,
                 size=10, color=WHITE_DIM)

    # exclusividade
    add_card(s, 770, 492, 446, 150, fill=RGBColor(0x33, 0x05, 0x0E))
    add_text(s, 790, 510, 350, 14, 'EXCLUSIVIDADE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 790, 535, 410, 100,
             'Estúdio próprio Spin no Brasil + dealer formado pela Spin '
             'Academy = a única operação capaz de entregar esse formato '
             'hoje no país.',
             size=10, color=WHITE_DIM, line_spacing=1.5)


# ---------------------------------------------------------------------- 06
def slide_06():
    s = new_slide(6, tagline='A PROVA · CASE THIAGO MASSAO',
                  caption='SPIN GAMING  ×  BLAZE  ·  CASE MASSAO')
    add_section_tabs(s, casting_active=True)

    add_eyebrow(s, 64, 130, 240, 'CASE REAL · MARÇO/2026')
    add_text(s, 64, 170, 600, 180,
             [('QUASE ', {}),
              ('1 MILHÃO', {'color': RED_BLAZE}),
              ('\nDE VIEWS COM\nUMA ÚNICA AÇÃO.', {})],
             size=36, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.05)
    add_text(s, 64, 360, 600, 100,
             ('O empresário e criador Thiago Massao foi treinado pela Spin '
              'Academy e assumiu, ao vivo, uma mesa de Baccarat operada pela '
              'Spin Gaming. A ação rodou em parceria com um cliente Spin e teve '
              'cobertura espontânea da Jovem Pan e do ConexãoBet, sendo descrita '
              'como "ação inédita no cenário brasileiro de iGaming".'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)
    add_text(s, 64, 475, 600, 60,
             ('O formato comprovou três coisas: (1) o público responde, '
              '(2) a imprensa amplifica gratuitamente, (3) a Spin tem o único '
              'pipeline pronto para escalar isso.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)

    # 3 stats
    stats = [
        ('~1M', 'VIEWS NA AÇÃO', 'Soma de live + cortes + cobertura editorial'),
        ('1ª', 'AÇÃO DO TIPO NO BRASIL', 'Influenciador atuando como dealer ao vivo'),
        ('60min', 'MESA AO VIVO', 'Interação direta criador ↔ jogador via chat'),
    ]
    x = 64
    for num, label, sub in stats:
        add_text(s, x, 580, 200, 60, num,
                 size=38, color=RED_SOFT, bold=True, font=FONT_BLACK)
        add_text(s, x, 645, 200, 14, label,
                 size=8, color=WHITE_MUTE, bold=True, spacing=2)
        add_text(s, x, 665, 200, 50, sub,
                 size=9, color=WHITE_DIM, line_spacing=1.4)
        x += 210

    # right image (use a Blaze-related photo)
    img = IMG_DIR / 'mesa-blaze-real.png'
    if img.exists():
        s.shapes.add_picture(str(img), px(700), px(140), px(516), px(560))


# ---------------------------------------------------------------------- 07
def slide_07():
    s = new_slide(7, tagline='DA AUDIÊNCIA AO VOLUME DE APOSTAS',
                  caption='SPIN GAMING  ×  BLAZE  ·  FUNIL')
    add_section_tabs(s, casting_active=True)

    add_eyebrow(s, 64, 130, 220, 'A MECÂNICA DO FUNIL')
    add_text(s, 64, 170, 720, 110,
             [('POR QUE CASTING BLAZE GERA\n', {}),
              ('+UAP · +FTD · +GGR', {'color': RED_BLAZE}),
              (' EM POUCO TEMPO.', {})],
             size=28, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    add_text(s, 800, 195, 416, 110,
             ('O influenciador-dealer corta etapas do funil tradicional de '
              'aquisição. Ele transforma audiência fria em jogador único '
              'autenticado (UAP) diretamente dentro do produto — e não em '
              'um landing externo.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)

    # 5 funnel cards
    cards = [
        ('ETAPA 01', 'ALCANCE', '600k+',
         'CCV médio diário do Kick BR. Streamers de cassino entregam picos de 20k–60k live viewers por sessão.'),
        ('ETAPA 02', 'CONFIANÇA', '63%',
         'da Gen Z confia mais em criadores do que em marcas. O dealer é o criador → confiança máxima na mesa.'),
        ('ETAPA 03', 'CONVERSÃO UAP', '0,5–5%',
         'de viewers convertem em FTD em campanhas de streaming (benchmark iGaming 2025/26).'),
        ('ETAPA 04', 'CPA', '-35%',
         'menor que o influencer marketing tradicional, segundo dados da Uberman Agency.'),
        ('ETAPA 05', 'VOLUME NAS MESAS SPIN', '+++',
         'Audiência aterrissa direto na mesa Spin × Blaze. GGR atribuído mensurável por código de mesa.'),
    ]
    x = 64
    w = (1280 - 64*2 - 4 * 14) / 5  # 5 cards com gap 14
    for i, (eyebrow, title, num, body) in enumerate(cards):
        special = (i == 4)
        add_card(s, x, 360, w, 280,
                 fill=RGBColor(0x33, 0x05, 0x0E) if special else BG_CARD)
        add_text(s, x + 16, 376, w - 32, 14, eyebrow,
                 size=8, color=RED_BLAZE, bold=True, spacing=3)
        add_text(s, x + 16, 396, w - 32, 32, title,
                 size=11, color=WHITE, bold=True)
        add_text(s, x + 16, 442, w - 32, 50, num,
                 size=28, color=RED_SOFT, bold=True, font=FONT_BLACK)
        add_text(s, x + 16, 500, w - 32, 130, body,
                 size=8.5, color=WHITE_DIM, line_spacing=1.45)
        x += w + 14

    add_text(s, 64, 680, 1150, 14,
             'FONTES:  StreamCharts (2026) · Kinser Content iGaming Report 2026 · Uberman Agency · TaDa Gaming Brasil · BrasilVegas',
             size=7, color=WHITE_MUTE, align='right')


# ---------------------------------------------------------------------- 08
def slide_08():
    s = new_slide(8, tagline='PILAR 02 · CASHBACK 10% MESAS SPIN',
                  caption='SPIN GAMING  ×  BLAZE  ·  CASHBACK')
    add_section_tabs(s, cashback_active=True)

    add_eyebrow(s, 64, 130, 240, 'PROPOSTA INICIAL · V1')
    add_text(s, 64, 170, 620, 220,
             [('10% DE VOLTA SOBRE\nAS PERDAS LÍQUIDAS,\nTODA SEMANA, NAS\n', {}),
              ('MESAS SPIN', {'color': RED_BLAZE}),
              (' DA BLAZE.', {})],
             size=30, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    add_text(s, 64, 410, 620, 120,
             ('Devolução automática de 10% das perdas líquidas semanais '
              'apenas para apostas nas mesas Spin Gaming dentro da Blaze. '
              'Sem rollover abusivo, com cap configurável por segmento e '
              'crédito em saldo real — em linha com o que KTO, Superbet, '
              'Vbet, Bet365 e 1xBet já praticam no Brasil.'),
             size=10.5, color=WHITE_DIM, line_spacing=1.5)

    # 4 mini-cards
    mini = [
        ('PERIODICIDADE', 'Semanal · janela seg→dom · crédito automático na 2ª-feira.'),
        ('ESCOPO', 'Apenas mesas Spin · concentra a base ativa no nosso lobby.'),
        ('FORMATO', '10% das perdas líquidas · saldo real · sem wagering pesado.'),
        ('SEGMENTAÇÃO', 'Trigger por queda de frequência · acelera ofertas ao churn.'),
    ]
    positions = [(64, 555), (332, 555), (64, 645), (332, 645)]
    for (mx, my), (title, body) in zip(positions, mini):
        add_card(s, mx, my, 260, 80)
        add_text(s, mx + 18, my + 14, 220, 14, title,
                 size=8, color=RED_BLAZE, bold=True, spacing=3)
        add_text(s, mx + 18, my + 34, 230, 42, body,
                 size=9, color=WHITE_DIM, line_spacing=1.4)

    # right photo
    img = IMG_DIR / 'cashback_growth.png'
    if img.exists():
        s.shapes.add_picture(str(img), px(700), px(140), px(516), px(560))
    # regulatorio pill
    add_pill(s, 720, 160, 160, 26, 'REGULATÓRIO BR · OK',
             fill=RED_BLAZE, text_color=WHITE)
    # conformidade caption
    add_card(s, 720, 600, 480, 80, fill=RGBColor(0x14, 0x03, 0x09))
    add_text(s, 740, 615, 440, 14, 'CONFORMIDADE',
             size=8, color=WHITE_MUTE, bold=True, spacing=2)
    add_text(s, 740, 635, 440, 42,
             'Cashback segue permitido pela SPA/MF (Lei 14.790/2023). '
             'Acompanhamos PL 1018/2026 e teremos plano de descontinuidade '
             'rápido caso a regra mude.',
             size=8.5, color=WHITE, line_spacing=1.4)


# ---------------------------------------------------------------------- 09
def slide_09():
    s = new_slide(9, tagline='CASHBACK · BENCHMARK DE MERCADO',
                  caption='SPIN GAMING  ×  BLAZE  ·  BENCHMARK')
    add_section_tabs(s, cashback_active=True)

    add_eyebrow(s, 64, 140, 200, 'POR QUE FUNCIONA')
    add_text(s, 64, 180, 700, 180,
             [('O CASHBACK É O PRINCIPAL\nMOTOR DE ', {}),
              ('RETENÇÃO', {'color': RED_BLAZE}),
              ('\nDAS BETS NO BRASIL.', {})],
             size=32, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    add_text(s, 820, 200, 396, 130,
             ('Com a proibição do bônus de boas-vindas em 2024, o cashback '
              'virou a principal arma legal de retenção das casas Tier-1 do '
              'país. Os efeitos são consistentes em todos os benchmarks do '
              'setor.'),
             size=11, color=WHITE_DIM, line_spacing=1.55)

    # 3 KPI cards
    kpis = [('+20–30%', 'RETENÇÃO 6 MESES'),
            ('+18%', 'FREQUÊNCIA DE DEPÓSITO'),
            ('+25%', 'TICKET MÉDIO')]
    x = 64
    w = (1280 - 64*2 - 20*2) / 3
    for num, label in kpis:
        add_card(s, x, 410, w, 130)
        add_text(s, x + 28, 430, w - 56, 70, num,
                 size=42, color=RED_SOFT, bold=True, font=FONT_BLACK)
        add_text(s, x + 28, 502, w - 56, 24, label,
                 size=10, color=WHITE_MUTE, bold=True, spacing=2)
        x += w + 20

    # Leitura
    add_card(s, 64, 580, 1152, 110, fill=RGBColor(0x2A, 0x07, 0x10))
    add_text(s, 90, 600, 1100, 14, 'LEITURA PARA A BLAZE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 90, 622, 1100, 70,
             [('Entrar com ', {}),
              ('10% semanal exclusivo nas mesas Spin', {'color': RED_SOFT, 'bold': True}),
              (' é um posicionamento competitivo, alinhado ao mercado e '
               'sustentável — sem inflar custo de bônus genérico. O recorte '
               'por mesa Spin permite medir ', {}),
              ('GGR incremental com precisão.', {'bold': True})],
             size=12, color=WHITE_DIM, line_spacing=1.55)


# ---------------------------------------------------------------------- 10
def slide_10():
    s = new_slide(10, tagline='CASHBACK · IMPACTO ESPERADO',
                  caption='SPIN GAMING  ×  BLAZE  ·  IMPACTO')
    add_section_tabs(s, cashback_active=True)

    add_eyebrow(s, 64, 140, 200, 'A LÓGICA')
    add_text(s, 64, 180, 600, 320,
             [('CASHBACK NÃO É CUSTO,\nÉ ', {}),
              ('AMORTECEDOR\nDE CHURN', {'color': RED_BLAZE}),
              (' DEPOIS\nDE UMA PERDA.', {})],
             size=42, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    # right narrative
    add_text(s, 720, 180, 496, 100,
             ('O momento mais perigoso na vida de um jogador é a primeira '
              'sessão perdedora. Sem nada que amenize a frustração, ele tende '
              'a abandonar a plataforma.'),
             size=12, color=WHITE_DIM, line_spacing=1.6)
    add_text(s, 720, 290, 496, 140,
             [('Ao receber ', {}),
              ('10% de volta na semana', {'bold': True}),
              (', o jogador volta para a plataforma para usar o crédito — e '
               'quando ele volta, encontra justamente as mesas Spin novamente. '
               'Cria-se um ', {}),
              ('ciclo de re-deposit fechado', {'color': RED_SOFT, 'bold': True}),
              (' dentro do nosso produto.', {})],
             size=12, color=WHITE_DIM, line_spacing=1.6)

    add_card(s, 720, 480, 496, 200, fill=RGBColor(0x33, 0x05, 0x0E))
    add_text(s, 740, 500, 460, 14, 'EFEITO COMPOSTO COM O CASTING BLAZE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 740, 525, 460, 140,
             [('O Casting Blaze ', {}),
              ('traz', {'bold': True}),
              (' jogadores novos para as mesas Spin. O Cashback ', {}),
              ('mantém', {'bold': True}),
              (' esses jogadores e maximiza re-deposits. As duas alavancas '
               'atacam pontas opostas do LTV: aquisição e retenção.', {})],
             size=11, color=WHITE_DIM, line_spacing=1.55)


# ---------------------------------------------------------------------- 11
def slide_11():
    s = new_slide(11, tagline='A SINERGIA · CASTING + CASHBACK',
                  caption='SPIN GAMING  ×  BLAZE  ·  PLANO')
    add_section_tabs(s, casting_active=True, cashback_active=True)

    add_eyebrow(s, 64, 140, 240, 'JUNTOS, NÃO SEPARADOS')
    add_text(s, 64, 180, 1150, 80,
             [('UMA AÇÃO FAZ O JOGADOR ENTRAR.\nA OUTRA FAZ ELE ', {}),
              ('FICAR — E VOLTAR.', {'color': RED_BLAZE})],
             size=26, color=WHITE, bold=True, font=FONT_BLACK, line_spacing=1.1)

    # 2 cards
    add_card(s, 64, 290, 565, 200)
    add_text(s, 88, 308, 530, 14, 'FUNIL DE AQUISIÇÃO · CASTING BLAZE',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 88, 330, 510, 24, 'TRAZ NOVOS UAP PARA AS MESAS SPIN',
             size=12, color=WHITE, bold=True)
    add_bullet_list(s, 88, 365, 510, [
        'Conteúdo orgânico do criador + mesa-código exclusiva',
        'Imprensa espontânea amplifica (case Massao ~1M views)',
        'CPA ~35% menor que tráfego pago tradicional',
        'Audiência aterrissa direto na mesa Spin × Blaze',
    ], item_h=22, size=10)

    add_card(s, 651, 290, 565, 200)
    add_text(s, 675, 308, 530, 14, 'CICLO DE RETENÇÃO · CASHBACK 10%',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    add_text(s, 675, 330, 510, 24, 'SEGURA E REATIVA QUEM JÁ ESTÁ DENTRO',
             size=12, color=WHITE, bold=True)
    add_bullet_list(s, 675, 365, 510, [
        'Reduz churn pós-derrota (momento crítico)',
        '+18% frequência de depósito documentada',
        'Devolução só vale nas mesas Spin → re-deposit fechado',
        'Trigger por queda de atividade → ataca churn antes de acontecer',
    ], item_h=22, size=10)

    # Timeline 4 phases
    add_card(s, 64, 520, 1152, 220)
    add_text(s, 88, 540, 600, 14, 'PLANO DE ATIVAÇÃO PROPOSTO',
             size=8, color=RED_BLAZE, bold=True, spacing=3)
    phases = [
        ('F1', 'SETUP',
         'Definição de mesa-código Spin × Blaze · seleção de 3 influenciadores piloto · onboarding Spin Academy.'),
        ('F2', 'SOFT LAUNCH',
         '1ª transmissão Casting Blaze · cashback 10% nas mesas Spin ativado em paralelo · monitoramento em D+1.'),
        ('F3', 'ESCALA',
         'Calendário fixo de Casting Blaze (mín. 2 lives/semana) · ofertas de cashback triggered por queda de frequência.'),
        ('F4', 'OTIMIZAÇÃO',
         'Calibragem do % de cashback por segmento de LTV · decisão de exclusividade Spin × Blaze para o formato no Brasil.'),
    ]
    x = 88
    w = (1152 - 24 * 2 - 18 * 3) / 4
    for label, title, body in phases:
        add_circle(s, x, 575, 28, fill=RED_BLAZE, text=label, text_size=9)
        add_text(s, x + 40, 580, w - 40, 16, title,
                 size=10, color=WHITE, bold=True, spacing=1)
        add_text(s, x, 615, w, 110, body,
                 size=9, color=WHITE_DIM, line_spacing=1.45)
        x += w + 18


# ---------------------------------------------------------------------- 12
def slide_12():
    s = new_slide(12, tagline='PRÓXIMOS PASSOS',
                  caption='SPIN GAMING  ×  BLAZE  ·  OBRIGADO')

    add_text(s, 0, 130, 1280, 24, 'VAMOS LIGAR O CASTING BLAZE',
             size=10, color=RED_BLAZE, bold=True, align='center', spacing=3)
    add_rect(s, 580, 158, 120, 1, fill=RED_BLAZE)

    add_text(s, 64, 185, 1152, 220,
             [('TRANSFORME A INFLUÊNCIA\nQUE A BLAZE JÁ TEM EM\n', {}),
              ('VOLUME REAL DE APOSTAS.', {'color': RED_BLAZE})],
             size=44, color=WHITE, bold=True, font=FONT_BLACK,
             align='center', line_spacing=1.1)

    add_text(s, 240, 410, 800, 80,
             [('A Spin Gaming entrega ', {}),
              ('estúdio', {'bold': True}), (', ', {}),
              ('tecnologia', {'bold': True}), (', ', {}),
              ('Spin Academy', {'bold': True}), (' e ', {}),
              ('operação', {'bold': True}),
              ('. A Blaze entrega ', {}),
              ('marca', {'bold': True}), (', ', {}),
              ('base', {'bold': True}), (' e ', {}),
              ('distribuição', {'bold': True}),
              ('. Juntas, lançamos o primeiro programa recorrente de '
               'influenciador-dealer + cashback do mercado brasileiro.', {})],
             size=12, color=WHITE_DIM, align='center', line_spacing=1.55)

    # 3 cards
    cards = [
        ('PRÓXIMA REUNIÃO', 'SHORTLIST DE INFLUENCIADORES',
         'Spin envia em até 5 dias úteis 8 nomes com fit Blaze, CCV comprovado '
         'e disponibilidade de agenda.', False),
        ('SETUP DE CASHBACK', 'REGRA & MENSURAÇÃO',
         'Spin + Blaze alinham a engine de cashback, cap por segmento e dashboard '
         'semanal de impacto no GGR.', False),
        ('SOFT LAUNCH', '1ª LIVE CASTING BLAZE',
         'Soft launch da operação combinada — Casting Blaze ao vivo + cashback '
         'ativo nas mesas Spin no mesmo dia.', True),
    ]
    x = 130
    cw = 340
    for eyebrow, title, body, hl in cards:
        fill = RGBColor(0x33, 0x05, 0x0E) if hl else BG_CARD
        add_card(s, x, 530, cw, 160, fill=fill)
        add_text(s, x + 22, 548, cw - 44, 14, eyebrow,
                 size=8, color=RED_BLAZE, bold=True, spacing=3)
        add_text(s, x + 22, 570, cw - 44, 28, title,
                 size=12, color=WHITE, bold=True, line_spacing=1.15)
        add_text(s, x + 22, 608, cw - 44, 80, body,
                 size=9, color=WHITE_DIM, line_spacing=1.5)
        x += cw + 10

    # Blaze x Spin lockup
    add_text(s, 500, 700, 110, 30, 'blaze',
             size=22, color=RED_BLAZE, bold=True, font=FONT_BLACK)
    add_text(s, 615, 700, 30, 30, '×',
             size=20, color=WHITE, bold=True, align='center')
    add_text(s, 650, 700, 130, 30, 'spin GAMING',
             size=18, color=RED_BLAZE, bold=True, italic=True, font=FONT_BLACK)


# ============================================================================
# Build
# ============================================================================
if __name__ == '__main__':
    print('Building editable PPTX...')
    for i, fn in enumerate([slide_01, slide_02, slide_03, slide_04, slide_05,
                            slide_06, slide_07, slide_08, slide_09, slide_10,
                            slide_11, slide_12], start=1):
        fn()
        print(f'  slide {i:02d} ok')
    prs.save(OUT)
    print(f'\nSaved: {OUT}')
    print(f'Size:  {OUT.stat().st_size / 1024 / 1024:.1f} MB')
