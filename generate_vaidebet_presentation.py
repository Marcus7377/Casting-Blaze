from __future__ import annotations

from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS_DIR = ROOT / "assets"
GENERATED_DIR = ASSETS_DIR / "generated"

MANUAL_P4 = ASSETS_DIR / "vaidebet_manual_p4.png"
LOGO_PAGE = ASSETS_DIR / "vaidebet_logo_p1.png"
UPLOAD_MANUAL = Path("/home/ubuntu/.cursor/projects/workspace/uploads/MINI_MANUAL_VAIDEBET_e2db.pdf")
UPLOAD_LOGO = Path("/home/ubuntu/.cursor/projects/workspace/uploads/VAIDEBET_LOGO_PDF_15bb.pdf")

PPT_OUTPUT = ROOT / "Apresentacao_Comercial_Spin_VaideBet.pptx"


def hex_to_rgb(value: str) -> RGBColor:
    value = value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


COLORS = {
    "black": hex_to_rgb("#000000"),
    "yellow": hex_to_rgb("#FEC537"),
    "white": hex_to_rgb("#FFFFFF"),
    "charcoal": hex_to_rgb("#0E1117"),
    "slate": hex_to_rgb("#1C2330"),
    "bluegray": hex_to_rgb("#253244"),
    "green": hex_to_rgb("#27D17F"),
}


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    if bold:
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    else:
        font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    return ImageFont.truetype(font_path, size=size)


def prepare_brand_assets() -> dict[str, Path]:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    logo_yellow_dark = GENERATED_DIR / "logo_yellow_dark.png"
    logo_black_yellow = GENERATED_DIR / "logo_black_yellow.png"
    logo_white_dark = GENERATED_DIR / "logo_white_dark.png"
    logo_single = GENERATED_DIR / "logo_single_black_yellow.png"

    source_manual = MANUAL_P4
    source_logo = LOGO_PAGE

    if not source_manual.exists() and UPLOAD_MANUAL.exists():
        try:
            import fitz  # type: ignore

            doc = fitz.open(str(UPLOAD_MANUAL))
            page = doc.load_page(3)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            source_manual = GENERATED_DIR / "_manual_source.png"
            pix.save(str(source_manual))
            doc.close()
        except Exception:
            source_manual = MANUAL_P4

    if not source_logo.exists() and UPLOAD_LOGO.exists():
        try:
            import fitz  # type: ignore

            doc = fitz.open(str(UPLOAD_LOGO))
            page = doc.load_page(0)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            source_logo = GENERATED_DIR / "_logo_source.png"
            pix.save(str(source_logo))
            doc.close()
        except Exception:
            source_logo = LOGO_PAGE

    if source_manual.exists():
        image = Image.open(source_manual).convert("RGB")
        w, h = image.size
        half_w, half_h = w // 2, h // 2
        image.crop((0, 0, half_w, half_h)).save(logo_yellow_dark)
        image.crop((half_w, 0, w, half_h)).save(logo_black_yellow)
        image.crop((half_w, half_h, w, h)).save(logo_white_dark)
    else:
        def make_logo(path: Path, bg: tuple[int, int, int], fg: tuple[int, int, int]) -> None:
            logo = Image.new("RGB", (842, 595), bg)
            draw = ImageDraw.Draw(logo)
            draw.text((145, 225), "VAIDEBET", fill=fg, font=load_font(120, True))
            logo.save(path)

        make_logo(logo_yellow_dark, (0, 0, 0), (254, 197, 55))
        make_logo(logo_black_yellow, (254, 197, 55), (0, 0, 0))
        make_logo(logo_white_dark, (0, 0, 0), (255, 255, 255))

    if source_logo.exists():
        Image.open(source_logo).convert("RGB").save(logo_single)
    else:
        Image.open(logo_black_yellow).save(logo_single)

    return {
        "logo_yellow_dark": logo_yellow_dark,
        "logo_black_yellow": logo_black_yellow,
        "logo_white_dark": logo_white_dark,
        "logo_single": logo_single,
    }


def draw_card(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    w: int,
    h: int,
    name: str,
    shift: str,
    status: str = "APROVADO",
) -> None:
    draw.rounded_rectangle(
        (x, y, x + w, y + h),
        radius=18,
        fill=(17, 23, 32),
        outline=(66, 79, 96),
        width=2,
    )
    draw.rectangle((x + 24, y + 26, x + w - 24, y + 170), fill=(46, 57, 72))
    draw.ellipse((x + 40, y + 45, x + 120, y + 125), fill=(254, 197, 55))
    draw.text((x + 140, y + 68), name, fill=(255, 255, 255), font=load_font(32, True))

    chip_w = 150
    draw.rounded_rectangle(
        (x + w - chip_w - 26, y + 30, x + w - 26, y + 76),
        radius=14,
        fill=(25, 158, 101),
    )
    draw.text((x + w - chip_w - 5, y + 43), status, fill=(9, 20, 32), font=load_font(18, True))

    draw.text((x + 30, y + 200), "MESA VIP", fill=(254, 197, 55), font=load_font(19, True))
    draw.text((x + 30, y + 238), f"TURNO: {shift}", fill=(206, 214, 224), font=load_font(20))
    draw.text((x + 30, y + 280), "BLACKJACK | ROLETA", fill=(163, 173, 186), font=load_font(18))


def create_mock_images() -> dict[str, Path]:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    dashboard = GENERATED_DIR / "casting_vaidebet_dashboard.png"
    profile = GENERATED_DIR / "casting_vaidebet_perfil.png"
    roteiro = GENERATED_DIR / "casting_vaidebet_roteiro.png"

    # Dashboard mock.
    image = Image.new("RGB", (1600, 900), (5, 9, 15))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 130), fill=(8, 14, 24))
    draw.text((40, 42), "VAIDEBET | CASTING DEALERS", fill=(254, 197, 55), font=load_font(42, True))

    stats = [("12", "DEALERS"), ("3", "JOGOS"), ("3", "TURNOS"), ("92%", "SLA")]
    sx = 40
    for value, label in stats:
        draw.rounded_rectangle((sx, 168, sx + 210, 288), radius=18, fill=(24, 32, 44), outline=(60, 73, 90), width=2)
        draw.text((sx + 28, 194), value, fill=(254, 197, 55), font=load_font(46, True))
        draw.text((sx + 30, 246), label, fill=(173, 183, 196), font=load_font(20))
        sx += 235

    filters = ["TODOS", "MANHA", "TARDE", "NOITE", "FEMININO", "MASCULINO"]
    fx = 40
    for name in filters:
        fill = (254, 197, 55) if name == "TARDE" else (25, 34, 47)
        txt = (0, 0, 0) if name == "TARDE" else (181, 192, 207)
        draw.rounded_rectangle((fx, 322, fx + 165, 370), radius=12, fill=fill)
        draw.text((fx + 36, 337), name, fill=txt, font=load_font(20, True))
        fx += 182

    draw.rounded_rectangle((1180, 322, 1560, 370), radius=12, fill=(25, 34, 47))
    draw.text((1210, 337), "Buscar dealer...", fill=(120, 134, 153), font=load_font(20))

    draw.text((40, 418), "ELENCO COMPLETO", fill=(235, 239, 245), font=load_font(26, True))
    draw.line((40, 452, 1560, 452), fill=(42, 54, 69), width=2)

    names = ["SARAH", "MARCOS", "LARISSA", "GUSTAVO"]
    shifts = ["TARDE", "NOITE", "TARDE", "MANHA"]
    for i, (name, shift) in enumerate(zip(names, shifts)):
        draw_card(draw, 40 + i * 390, 480, 360, 360, name, shift)

    image.save(dashboard)

    # Dealer profile mock.
    img2 = Image.new("RGB", (1400, 820), (6, 10, 16))
    d2 = ImageDraw.Draw(img2)
    d2.rounded_rectangle((20, 20, 1380, 800), radius=22, fill=(7, 13, 22), outline=(52, 68, 87), width=2)
    d2.rectangle((20, 20, 1380, 145), fill=(9, 16, 28))
    d2.text((46, 52), "PERFIL DO DEALER", fill=(254, 197, 55), font=load_font(36, True))
    d2.rectangle((48, 182, 250, 382), fill=(37, 49, 64))
    d2.ellipse((86, 225, 212, 351), fill=(254, 197, 55))
    d2.text((294, 196), "SARAH", fill=(248, 250, 252), font=load_font(64, True))
    d2.text((296, 274), "ANA CLARA", fill=(168, 179, 194), font=load_font(28))

    chips = ["BLACKJACK", "ROLETA", "BACCARAT", "FEMININO"]
    cx = 292
    for chip in chips:
        cw = 200 if chip == "BLACKJACK" else 160
        d2.rounded_rectangle((cx, 328, cx + cw, 374), radius=16, fill=(24, 33, 46))
        d2.text((cx + 18, 341), chip, fill=(254, 197, 55), font=load_font(20, True))
        cx += cw + 14

    d2.text((48, 438), "INFORMACOES", fill=(254, 197, 55), font=load_font(28, True))
    d2.rounded_rectangle((48, 470, 458, 580), radius=14, fill=(24, 33, 46))
    d2.rounded_rectangle((484, 470, 894, 580), radius=14, fill=(24, 33, 46))
    d2.text((74, 500), "TURNO", fill=(151, 163, 180), font=load_font(22))
    d2.text((74, 534), "TARDE", fill=(255, 255, 255), font=load_font(28, True))
    d2.text((510, 500), "NICKNAME", fill=(151, 163, 180), font=load_font(22))
    d2.text((510, 534), "SARAH", fill=(255, 255, 255), font=load_font(28, True))

    d2.text((48, 628), "PONTOS FORTES", fill=(254, 197, 55), font=load_font(28, True))
    d2.text(
        (48, 674),
        "Excelente postura e tecnica de Blackjack. Lider de engajamento\nnas lives de mesa dedicada da Vai de Bet.",
        fill=(206, 214, 224),
        font=load_font(26),
    )
    img2.save(profile)

    # Roteiro mock.
    img3 = Image.new("RGB", (1600, 900), (7, 10, 17))
    d3 = ImageDraw.Draw(img3)
    d3.text((56, 56), "ROTEIRO DE MESA - VAI DE BET", fill=(254, 197, 55), font=load_font(54, True))
    d3.text((58, 132), "ORIENTACOES DE CONDUCAO | SPIN GAMING", fill=(153, 166, 184), font=load_font(25))

    sections = [
        ("ABERTURA", [
            '"Ola jogadores, meu nome e [Nome] e estamos ao vivo na Vai de Bet!"',
            '"Escolha seus lugares e aproveite a rodada com responsabilidade."',
            '"Hoje temos promocao especial para a comunidade Vai de Bet."',
        ]),
        ("DURANTE O JOGO", [
            "Priorizar boas praticas de jogo e reforcar transparencia da mesa.",
            "Conectar promocoes ativas com clareza e linguagem da marca.",
            "Acionar suporte para comportamentos fora do roteiro.",
        ]),
        ("FECHAMENTO", [
            '"Obrigado por jogar conosco! Voltamos em instantes."',
            '"Parabens aos vencedores e boa sorte na proxima rodada!"',
            '"Segue a gente na plataforma para acompanhar as proximas lives."',
        ]),
    ]

    box_w = 500
    for i, (title, lines) in enumerate(sections):
        x = 56 + i * (box_w + 20)
        d3.rounded_rectangle((x, 200, x + box_w, 820), radius=16, fill=(21, 29, 41), outline=(60, 74, 93), width=2)
        d3.text((x + 26, 236), title, fill=(254, 197, 55), font=load_font(28, True))
        y = 300
        for line in lines:
            wrapped = textwrap.fill(line, width=36)
            d3.multiline_text((x + 26, y), wrapped, fill=(218, 225, 235), font=load_font(20), spacing=10)
            y += 122

    img3.save(roteiro)

    return {"dashboard": dashboard, "profile": profile, "roteiro": roteiro}


def add_full_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    bold: bool = False,
    color: RGBColor = COLORS["white"],
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"


def add_title_band(slide, title: str, subtitle: str) -> None:
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.78)
    ).fill.solid()
    shape = slide.shapes[-1]
    shape.fill.fore_color.rgb = COLORS["yellow"]
    shape.line.fill.background()
    add_text(slide, title, 0.45, 0.15, 8.8, 0.38, 18, True, COLORS["black"])
    add_text(slide, subtitle, 9.05, 0.15, 3.85, 0.38, 12, True, COLORS["black"], PP_ALIGN.RIGHT)


def build_presentation(brand_assets: dict[str, Path], mocks: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1 - Capa.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "PROPOSTA COMERCIAL | SPIN GAMING", "JUNHO 2026")
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(5.9), Inches(13.333), Inches(1.6)
    )
    footer = slide.shapes[-1]
    footer.fill.solid()
    footer.fill.fore_color.rgb = COLORS["yellow"]
    footer.line.fill.background()
    add_text(slide, "MESAS DEDICADAS PARA VAI DE BET", 0.7, 2.0, 8.8, 1.2, 48, True, COLORS["yellow"])
    add_text(
        slide,
        "Projeto exclusivo para acelerar GGR, engajamento e reconhecimento de marca no Live Casino.",
        0.7,
        3.4,
        9.2,
        1.0,
        20,
        False,
        COLORS["white"],
    )
    add_text(
        slide,
        "JOGUE COM RESPONSABILIDADE AUTORIZADA PELA PORTARIA SPA/MF N. 668",
        0.7,
        6.45,
        11.8,
        0.5,
        14,
        True,
        COLORS["black"],
    )
    slide.shapes.add_picture(str(brand_assets["logo_black_yellow"]), Inches(8.8), Inches(1.5), Inches(4.0), Inches(2.1))

    # Slide 2 - Contexto executivo.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["charcoal"])
    add_title_band(slide, "DESAFIO E OPORTUNIDADE", "VAI DE BET + SPIN")
    add_text(
        slide,
        "Por que este projeto e estrategico para a Vai de Bet",
        0.65,
        1.0,
        9.2,
        0.8,
        34,
        True,
        COLORS["yellow"],
    )
    bullets = [
        "Criar mesas exclusivas com identidade da Vai de Bet e operacao 24/7.",
        "Transformar dealers em canal de comunicacao e reforco de campanhas.",
        "Controlar resultado de ponta a ponta com dashboards diarios, semanais e mensais.",
        "Aumentar proximidade entre operadora e provedora com monitoria presencial.",
    ]
    y = 2.0
    for line in bullets:
        add_text(slide, f"- {line}", 0.9, y, 12.0, 0.65, 22, False, COLORS["white"])
        y += 0.85
    slide.shapes.add_picture(str(brand_assets["logo_yellow_dark"]), Inches(8.8), Inches(4.8), Inches(4.0), Inches(2.0))

    # Slide 3 - Diferenciais.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "A PROPOSTA DE VALOR SPIN", "5 PILARES COMPETITIVOS")
    add_text(slide, "A melhor opcao em mesas dedicadas", 0.65, 1.0, 7.4, 0.7, 34, True, COLORS["yellow"])
    pillars = [
        "Dealers Dedicados",
        "Casting Dealers",
        "Dashboard KPIs",
        "Dashboard Influenciadores",
        "Proximidade sem igual",
    ]
    descriptions = [
        "Time exclusivo para a Vai de Bet, com rotina, escala e rotacao por mesa.",
        "Gestao de elenco com performance mensal, ajustes de roteiro e promocoes.",
        "Acompanhamento completo de GGR, Turnover, UAP, Margem e ARPU.",
        "Leitura de campanha por live: agenda, FTDs, volume, canal e views.",
        "Espaco fisico aberto para operacao diaria, eventos e relacionamento premium.",
    ]
    x = 0.5
    for title, desc in zip(pillars, descriptions):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.5), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["slate"]
        card.line.color.rgb = COLORS["yellow"]
        add_text(slide, title, x + 0.15, 2.2, 2.2, 0.9, 18, True, COLORS["yellow"], PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.18, 3.0, 2.16, 3.4, 14, False, COLORS["white"], PP_ALIGN.LEFT)
        x += 2.58

    # Slide 4 - Dealers Dedicados.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["charcoal"])
    add_title_band(slide, "1. DEALERS DEDICADOS", "OPERACAO EXCLUSIVA VAI DE BET")
    add_text(
        slide,
        "Dealers que atuam somente na Vai de Bet",
        0.7,
        1.1,
        7.0,
        0.8,
        32,
        True,
        COLORS["yellow"],
    )
    ops = [
        "Escala personalizada por turno e por jogo.",
        "Rotacao de mesas para manter energia e consistencia.",
        "Treinamento orientado para tom de voz da marca.",
        "Roteiro ativo com foco em engajamento e seguranca.",
    ]
    y = 2.0
    for item in ops:
        add_text(slide, f"- {item}", 0.9, y, 7.0, 0.55, 20, False, COLORS["white"])
        y += 0.72
    slide.shapes.add_picture(str(mocks["profile"]), Inches(7.4), Inches(1.3), Inches(5.5), Inches(5.8))

    # Slide 5 - Casting dashboard.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "2. CASTING DEALERS", "DASHBOARD DE ELENCO")
    add_text(slide, "Gestao de elenco e performance em um unico painel", 0.7, 1.0, 10.8, 0.7, 30, True, COLORS["yellow"])
    slide.shapes.add_picture(str(mocks["dashboard"]), Inches(0.7), Inches(1.8), Inches(12.0), Inches(4.9))
    add_text(
        slide,
        "Filtros por turno, jogo e perfil + indicadores para acao rapida do time comercial.",
        0.8,
        6.9,
        12.0,
        0.4,
        15,
        False,
        COLORS["white"],
    )

    # Slide 6 - Roteiro e orientacoes.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["charcoal"])
    add_title_band(slide, "3. ROTEIRO PERSONALIZADO", "COMUNICACAO CONSISTENTE")
    add_text(slide, "Dealer como canal de comunicacao da marca", 0.7, 1.0, 9.0, 0.7, 30, True, COLORS["yellow"])
    add_text(
        slide,
        "Atualize aberturas, promocoes e scripts de mesa em tempo real para manter a campanha alinhada.",
        0.8,
        1.7,
        8.8,
        1.0,
        18,
        False,
        COLORS["white"],
    )
    slide.shapes.add_picture(str(mocks["roteiro"]), Inches(0.7), Inches(2.6), Inches(12.0), Inches(4.6))

    # Slide 7 - KPIs.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "4. DASHBOARD KPIS", "GESTAO DIARIA, SEMANAL E MENSAL")
    add_text(slide, "Inteligencia de resultado para tomada de decisao", 0.7, 1.0, 10.8, 0.7, 30, True, COLORS["yellow"])

    kpis = ["GGR", "TURNOVER", "UAP", "MARGEM", "QTD APOSTAS", "MEDIA APOSTAS", "ARPU"]
    positions = [(0.8, 2.0), (2.6, 2.0), (4.4, 2.0), (6.2, 2.0), (8.0, 2.0), (9.8, 2.0), (11.6, 2.0)]
    for (x, y), metric in zip(positions, kpis):
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.6), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = COLORS["bluegray"]
        box.line.color.rgb = COLORS["yellow"]
        add_text(slide, metric, x + 0.08, y + 0.17, 1.44, 0.3, 12, True, COLORS["yellow"], PP_ALIGN.CENTER)
        add_text(slide, "Acompanhar", x + 0.08, y + 0.53, 1.44, 0.3, 12, False, COLORS["white"], PP_ALIGN.CENTER)

    add_text(
        slide,
        "Com a leitura consolidada, a Spin e a Vai de Bet ajustam escala, roteiro e ativacao para acelerar crescimento com eficiencia.",
        0.8,
        3.5,
        11.8,
        1.2,
        19,
        False,
        COLORS["white"],
    )
    add_text(slide, "Ciclos de report:", 0.8, 5.0, 2.0, 0.4, 16, True, COLORS["yellow"])
    add_text(slide, "Diario | Semanal | Mensal", 2.3, 5.0, 4.2, 0.4, 16, True, COLORS["white"])

    timeline = [
        ("DIA 1", "Leitura de base e comportamento"),
        ("DIA 7", "Ajuste fino de turno e promocao"),
        ("DIA 30", "Plano de escala para ganho de margem"),
    ]
    tx = 0.8
    for label, text in timeline:
        marker = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(tx), Inches(5.8), Inches(3.95), Inches(1.2))
        marker.fill.solid()
        marker.fill.fore_color.rgb = COLORS["slate"]
        marker.line.color.rgb = COLORS["yellow"]
        add_text(slide, label, tx + 0.15, 6.0, 3.6, 0.3, 16, True, COLORS["yellow"], PP_ALIGN.CENTER)
        add_text(slide, text, tx + 0.2, 6.3, 3.5, 0.5, 13, False, COLORS["white"], PP_ALIGN.CENTER)
        tx += 4.15

    # Slide 8 - Influenciadores.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["charcoal"])
    add_title_band(slide, "5. DASHBOARD DE INFLUENCIADORES", "PERFORMANCE DE CAMPANHAS")
    add_text(slide, "Campanhas de live com leitura completa de conversao", 0.7, 1.0, 10.8, 0.7, 30, True, COLORS["yellow"])
    metrics = ["AGENDA DE LIVES", "PERFORMANCE DA LIVE", "FTDS", "VOLUME FINANCEIRO", "CANAL", "NUMERO DE VIEWS"]
    x = 0.8
    y = 2.0
    for i, metric in enumerate(metrics):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(4.0), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["slate"]
        card.line.color.rgb = COLORS["yellow"]
        add_text(slide, metric, x + 0.1, y + 0.2, 3.8, 0.3, 14, True, COLORS["yellow"], PP_ALIGN.CENTER)
        add_text(slide, "Indicador estrategico", x + 0.1, y + 0.58, 3.8, 0.3, 13, False, COLORS["white"], PP_ALIGN.CENTER)
        if i % 3 == 2:
            x = 0.8
            y += 1.35
        else:
            x += 4.2
    add_text(
        slide,
        "Visao integrada para decidir quais creators escalar, quais formatos repetir e como converter audiencia em deposito.",
        0.8,
        5.1,
        12.0,
        1.2,
        18,
        False,
        COLORS["white"],
    )
    growth = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(9.6), Inches(5.6), Inches(3.0), Inches(1.2))
    growth.fill.solid()
    growth.fill.fore_color.rgb = COLORS["yellow"]
    growth.line.fill.background()
    add_text(slide, "ESCALA DE ROI", 9.9, 5.95, 2.3, 0.3, 15, True, COLORS["black"], PP_ALIGN.CENTER)

    # Slide 9 - Proximidade.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "PROXIMIDADE OPERACIONAL", "ESTRUTURA ABERTA PARA A OPERADORA")
    add_text(slide, "Presenca diaria da Vai de Bet dentro da operacao", 0.7, 1.0, 11.2, 0.7, 32, True, COLORS["yellow"])
    topics = [
        "Acompanhamento in loco da rotina de mesas e dealers.",
        "Correcao rapida de oportunidades e reforco de boas praticas.",
        "Espaco para eventos de marca: pre-lancamento, pos-lancamento e high rollers.",
        "Comite conjunto Spin + Vai de Bet para evolucao continua.",
    ]
    y = 2.0
    for item in topics:
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(12.0), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["slate"]
        card.line.color.rgb = COLORS["yellow"]
        add_text(slide, item, 1.0, y + 0.24, 11.6, 0.4, 18, False, COLORS["white"])
        y += 1.1
    add_text(slide, "MODELO DE RELACIONAMENTO: SPIN COMO EXTENSAO DO TIME VAI DE BET", 0.8, 6.5, 12.0, 0.5, 16, True, COLORS["yellow"], PP_ALIGN.CENTER)

    # Slide 10 - Implementacao.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["charcoal"])
    add_title_band(slide, "IMPLEMENTACAO", "GO LIVE EM ATE 60 DIAS")
    add_text(slide, "Plano de execucao da parceria", 0.7, 1.0, 7.0, 0.8, 33, True, COLORS["yellow"])

    phases = [
        ("1. Kickoff", "Definicao de metas, escopo e prioridades da Vai de Bet."),
        ("2. Design e Branding", "Ajuste cenografico, scripts e identidade visual das mesas."),
        ("3. Selecao de Dealers", "Casting, treinamento e validacao do elenco dedicado."),
        ("4. Go Live + Otimizacao", "Inicio da operacao com monitoria continua de KPIs."),
    ]
    y = 1.9
    for title, desc in phases:
        bullet = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y), Inches(12.0), Inches(1.0))
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = COLORS["slate"]
        bullet.line.color.rgb = COLORS["yellow"]
        add_text(slide, title, 1.05, y + 0.14, 3.0, 0.3, 17, True, COLORS["yellow"])
        add_text(slide, desc, 3.3, y + 0.15, 9.0, 0.6, 16, False, COLORS["white"])
        y += 1.2

    badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(6.3), Inches(3.2), Inches(0.8))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLORS["yellow"]
    badge.line.fill.background()
    add_text(slide, "OPERACAO 24/7", 9.7, 6.52, 2.8, 0.3, 16, True, COLORS["black"], PP_ALIGN.CENTER)

    # Slide 11 - Proximos passos.
    slide = prs.slides.add_slide(blank)
    add_full_background(slide, COLORS["black"])
    add_title_band(slide, "PROXIMOS PASSOS", "VAI DE BET + SPIN")
    add_text(slide, "Vamos colocar a mesa dedicada no ar", 0.7, 1.0, 8.4, 0.8, 36, True, COLORS["yellow"])
    steps = [
        "Validacao comercial da proposta.",
        "Ajustes finais de branding e roteiro.",
        "Definicao conjunta do cronograma de implementacao.",
        "Kickoff operacional e inicio das ativacoes.",
    ]
    x = 0.8
    for step in steps:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.0), Inches(3.05), Inches(2.3))
        box.fill.solid()
        box.fill.fore_color.rgb = COLORS["slate"]
        box.line.color.rgb = COLORS["yellow"]
        add_text(slide, step, x + 0.17, 3.55, 2.7, 1.5, 16, True, COLORS["white"], PP_ALIGN.CENTER)
        x += 3.2

    slide.shapes.add_picture(str(brand_assets["logo_white_dark"]), Inches(4.8), Inches(5.8), Inches(3.8), Inches(1.6))
    add_text(slide, "Obrigado.", 5.8, 6.95, 2.0, 0.35, 20, True, COLORS["yellow"], PP_ALIGN.CENTER)

    prs.save(PPT_OUTPUT)


def main() -> None:
    brand_assets = prepare_brand_assets()
    mocks = create_mock_images()
    build_presentation(brand_assets, mocks)
    print(f"Presentation generated: {PPT_OUTPUT}")


if __name__ == "__main__":
    main()
