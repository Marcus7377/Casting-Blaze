from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)

COL_BG = RGBColor(24, 4, 26)
COL_BG_ACCENT = RGBColor(95, 0, 44)
COL_PRIMARY = RGBColor(255, 45, 102)
COL_PRIMARY_SOFT = RGBColor(199, 59, 104)
COL_TEXT = RGBColor(245, 245, 245)
COL_MUTED = RGBColor(209, 193, 201)


def add_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COL_BG

    glow = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, HEIGHT)
    glow.fill.solid()
    glow.fill.fore_color.rgb = COL_BG_ACCENT
    glow.fill.transparency = 72
    glow.line.fill.background()

    side = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.5), 0, Inches(2.9), HEIGHT)
    side.fill.solid()
    side.fill.fore_color.rgb = RGBColor(150, 10, 65)
    side.fill.transparency = 83
    side.line.fill.background()

    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.15), Inches(0.15), Inches(13.03), Inches(7.2))
    border.fill.background()
    border.line.color.rgb = RGBColor(111, 49, 79)
    border.line.width = Pt(1.1)


def write_text(slide, left, top, width, height, text, size=20, bold=False, color=COL_TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    font = p.font
    font.name = "Montserrat"
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    return box


def write_bullets(slide, left, top, width, height, bullets, size=18, color=COL_TEXT, level=0, spacing=1.2):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = level
        p.line_spacing = spacing
        p.space_after = Pt(8)
        font = p.font
        font.name = "Montserrat"
        font.size = Pt(size)
        font.color.rgb = color
    return box


def add_header(slide, kicker, title, subtitle=None):
    write_text(
        slide,
        Inches(0.55),
        Inches(0.35),
        Inches(8.8),
        Inches(0.4),
        kicker.upper(),
        size=12,
        bold=True,
        color=COL_PRIMARY,
    )
    write_text(
        slide,
        Inches(0.55),
        Inches(0.72),
        Inches(10.8),
        Inches(1.15),
        title,
        size=38,
        bold=True,
    )
    if subtitle:
        write_text(
            slide,
            Inches(0.55),
            Inches(1.78),
            Inches(10.8),
            Inches(0.8),
            subtitle,
            size=17,
            color=COL_MUTED,
        )

    write_text(
        slide,
        Inches(10.05),
        Inches(0.4),
        Inches(2.9),
        Inches(0.4),
        "SPIN GAMING  x  BLAZE",
        size=11,
        bold=True,
        color=COL_TEXT,
        align=PP_ALIGN.RIGHT,
    )


def card(slide, left, top, width, height, title, lines):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(49, 10, 35)
    shape.fill.transparency = 20
    shape.line.color.rgb = COL_PRIMARY_SOFT
    shape.line.width = Pt(1.3)

    write_text(
        slide,
        left + Inches(0.2),
        top + Inches(0.18),
        width - Inches(0.4),
        Inches(0.5),
        title,
        size=15,
        bold=True,
        color=COL_PRIMARY,
    )
    write_bullets(
        slide,
        left + Inches(0.24),
        top + Inches(0.67),
        width - Inches(0.48),
        height - Inches(0.85),
        lines,
        size=13,
        color=COL_TEXT,
        spacing=1.15,
    )


def media_placeholder(slide, left, top, width, height, title, subtitle):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(30, 8, 25)
    frame.fill.transparency = 8
    frame.line.color.rgb = COL_PRIMARY
    frame.line.width = Pt(1.8)

    icon = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        left + (width / 2) - Inches(0.25),
        top + Inches(0.6),
        Inches(0.5),
        Inches(0.42),
    )
    icon.fill.solid()
    icon.fill.fore_color.rgb = COL_PRIMARY
    icon.line.fill.background()
    icon.rotation = 90

    write_text(
        slide,
        left + Inches(0.18),
        top + Inches(1.15),
        width - Inches(0.36),
        Inches(0.45),
        title,
        size=15,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    write_text(
        slide,
        left + Inches(0.18),
        top + Inches(1.57),
        width - Inches(0.36),
        Inches(0.7),
        subtitle,
        size=11,
        color=COL_MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_divider(slide, top):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), top, Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(112, 56, 83)
    line.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = WIDTH
    prs.slide_height = HEIGHT

    # Slide 1 - Capa
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Proposta Estrategica  |  Junho 2026",
        "Casting de Influencer Dealer para Blaze",
        "Como converter alcance em UAPI, depositos e maior volume de apostas nas mesas da Spin Gaming.",
    )
    card(
        s,
        Inches(0.55),
        Inches(2.6),
        Inches(6.15),
        Inches(2.8),
        "Objetivo da proposta",
        [
            "Escalar para Blaze o formato de influenciador atuando como dealer ao vivo.",
            "Replicar e superar o case de quase 1 milhao de views ja validado pela Spin.",
            "Acoplar campanha de cashback de 10% para acelerar reativacao e recorrencia.",
        ],
    )
    card(
        s,
        Inches(6.95),
        Inches(2.6),
        Inches(5.8),
        Inches(2.8),
        "Tese central",
        [
            "Influencer dealer aumenta autenticidade percebida e tempo de sessao.",
            "Cashback reduz aversao a perda e mantem jogador ativo na jornada.",
            "Combinacao eleva UAPI, depositares recorrentes e turnover de mesa.",
        ],
    )
    write_text(
        s,
        Inches(0.55),
        Inches(6.92),
        Inches(12.0),
        Inches(0.3),
        "Documento de trabalho comercial | Spin Gaming",
        size=10,
        color=COL_MUTED,
    )

    # Slide 2 - Contexto de mercado
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Contexto de mercado",
        "O Brasil entrou na fase de consolidacao e guerra por retencao.",
    )
    write_bullets(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(6.4),
        Inches(4.5),
        [
            "25,2 milhoes de apostadores unicos (CPF) no mercado regulado em 2025.",
            "R$37 bi de GGR em 2025 no primeiro ano completo de regulacao.",
            "Mais de 100 milhoes de contas registradas, indicando forte multihoming.",
            "Em mercados competitivos, diferencial de produto e experiencia passa a decidir share de wallet.",
        ],
        size=20,
    )
    card(
        s,
        Inches(7.4),
        Inches(2.05),
        Inches(5.2),
        Inches(3.5),
        "Implicacao para a Blaze",
        [
            "Nao basta aquisicao: precisa criar razao recorrente para voltar.",
            "Live Casino com creator conhecido gera memorabilidade da marca.",
            "Promocao inteligente (cashback) protege margem e sustenta frequencia.",
        ],
    )
    write_text(
        s,
        Inches(0.7),
        Inches(6.85),
        Inches(12.2),
        Inches(0.35),
        "Fontes: SPA/MF (2026), InterGame (jan/2026), SBC Noticias (jan/2026).",
        size=10,
        color=COL_MUTED,
    )

    # Slide 3 - Prova de conceito
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Prova de conceito Spin",
        "A acao com influenciador dealer ja mostrou tracao real.",
    )
    media_placeholder(
        s,
        Inches(0.75),
        Inches(2.0),
        Inches(4.0),
        Inches(2.15),
        "Video 1 - Massao como dealer",
        "Inserir video anexado pela equipe Spin",
    )
    media_placeholder(
        s,
        Inches(4.95),
        Inches(2.0),
        Inches(4.0),
        Inches(2.15),
        "Video 2 - Massao como dealer",
        "Inserir video anexado pela equipe Spin",
    )
    media_placeholder(
        s,
        Inches(9.15),
        Inches(2.0),
        Inches(3.45),
        Inches(2.15),
        "Print do case",
        "Acao com quase 1 milhao de visualizacoes",
    )
    add_divider(s, Inches(4.45))
    write_bullets(
        s,
        Inches(0.85),
        Inches(4.7),
        Inches(12.0),
        Inches(1.9),
        [
            "Quase 1 milhao de views confirma potencial de alcance + conversa social.",
            "Formato aproxima entretenimento e experiencia de mesa real.",
            "Para Blaze, o ganho vem da conversao desse alcance em jogadores autenticos ativos (UAPI).",
        ],
        size=17,
        spacing=1.12,
    )

    # Slide 4 - Por que funciona
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Motor de crescimento",
        "Por que o influenciador dealer aumenta UAPI e depositos.",
    )
    card(
        s,
        Inches(0.65),
        Inches(2.0),
        Inches(4.0),
        Inches(3.8),
        "1) Credibilidade transferida",
        [
            "A comunidade ja confia no creator.",
            "Essa confianca reduz friccao de primeiro deposito.",
            "Jogador entra por afinidade, nao apenas por bonus.",
        ],
    )
    card(
        s,
        Inches(4.85),
        Inches(2.0),
        Inches(4.0),
        Inches(3.8),
        "2) Conteudo com prova social",
        [
            "Clipes da mesa geram FOMO e comentario em tempo real.",
            "Momentos de ganho e reacao elevam taxa de compartilhamento.",
            "Cada sessao alimenta novo topo de funil organico.",
        ],
    )
    card(
        s,
        Inches(9.05),
        Inches(2.0),
        Inches(3.65),
        Inches(3.8),
        "3) Jornada mais longa",
        [
            "Live aumenta tempo de permanencia.",
            "Sessao mais longa costuma elevar total apostado.",
            "Recorrencia cria base de apostador autentico.",
        ],
    )

    # Slide 5 - Arquitetura UAPI
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Arquitetura de resultado",
        "Modelo UAPI: do alcance qualificado ao volume de apostas.",
    )
    write_text(s, Inches(0.8), Inches(2.0), Inches(12.0), Inches(0.4), "FUNIL PROPOSTO", size=14, bold=True, color=COL_PRIMARY)

    steps = [
        ("ALCANCE QUALIFICADO", "Audiencia do influencer com afinidade em jogo."),
        ("UAPI", "Entrada de jogadores autenticos com intencao real de jogar."),
        ("1o DEPOSITO", "Conversao em deposito inicial durante janela quente de interesse."),
        ("RECORRENCIA", "Retorno em multiplas sessoes por conexao com o formato."),
        ("VOLUME", "Aumento do turnover total das mesas Spin na Blaze."),
    ]

    start_left = Inches(0.8)
    top = Inches(2.5)
    w = Inches(2.35)
    gap = Inches(0.25)
    for idx, (ttl, desc) in enumerate(steps):
        l = start_left + idx * (w + gap)
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, top, w, Inches(2.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(42, 10, 34)
        shape.line.color.rgb = COL_PRIMARY_SOFT
        shape.line.width = Pt(1.2)
        write_text(s, l + Inches(0.12), top + Inches(0.12), w - Inches(0.24), Inches(0.52), ttl, size=12, bold=True, color=COL_PRIMARY, align=PP_ALIGN.CENTER)
        write_text(s, l + Inches(0.14), top + Inches(0.7), w - Inches(0.28), Inches(1.8), desc, size=12, color=COL_TEXT, align=PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l + w + Inches(0.02), top + Inches(1.05), Inches(0.2), Inches(0.5))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = COL_PRIMARY
            arrow.line.fill.background()

    write_text(
        s,
        Inches(0.8),
        Inches(6.1),
        Inches(12.0),
        Inches(0.8),
        "Metrica principal sugerida: UAPI 30d = novos jogadores com 2+ sessoes e 2+ depositos em ate 30 dias.",
        size=13,
        color=COL_MUTED,
    )

    # Slide 6 - Formato da acao
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Formato proposto para Blaze",
        "Casting de influencer dealer em temporadas curtas e escalaveis.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(5.95),
        Inches(3.7),
        "Design da temporada piloto (8 semanas)",
        [
            "2 lives por semana com influencer dealer (90 min por sessao).",
            "Quadro fixo de agenda: Dia A (aquecimento) + Dia B (mesa principal).",
            "Drops de conteudo curto (reels/cortes) em ate 24h apos cada live.",
            "Landing de campanha com CTA direto para mesas Spin na Blaze.",
        ],
    )
    card(
        s,
        Inches(6.9),
        Inches(2.0),
        Inches(5.7),
        Inches(3.7),
        "Governanca de marca e compliance",
        [
            "Roteiro aprovado previamente por Blaze + Spin.",
            "Mensagens claras de jogo responsavel e 18+ em todas as pecas.",
            "Moderacao de chat e politica de conduta durante as transmissoes.",
            "Modelo replicavel para novos creators apos validacao inicial.",
        ],
    )

    # Slide 7 - KPI e metas
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Metas e mensuracao",
        "KPI framework para provar impacto em negocio.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(3.9),
        Inches(3.8),
        "Topo de funil",
        [
            "Views qualificadas dos conteudos",
            "CTR para landing da campanha",
            "Taxa de cadastro originada do creator",
        ],
    )
    card(
        s,
        Inches(4.85),
        Inches(2.0),
        Inches(3.9),
        Inches(3.8),
        "Conversao",
        [
            "1o deposito (FTD) vindo da campanha",
            "Custo por FTD autenticado",
            "UAPI 30d e UAPI 60d por coorte",
        ],
    )
    card(
        s,
        Inches(9.0),
        Inches(2.0),
        Inches(3.6),
        Inches(3.8),
        "Monetizacao",
        [
            "Volume apostado por mesa",
            "Depositos por jogador ativo",
            "NGR incremental apos custo promo",
        ],
    )
    write_text(
        s,
        Inches(0.7),
        Inches(6.85),
        Inches(12.0),
        Inches(0.35),
        "Recomendacao: painel semanal com leitura por coorte e comparativo versus baseline de mesas.",
        size=10,
        color=COL_MUTED,
    )

    # Slide 8 - Cashback overview
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Programa de cashback para mesas Spin",
        "Oferta inicial: 10% de cashback para perdas liquidas elegiveis.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(6.0),
        Inches(3.9),
        "Mecanica sugerida",
        [
            "Cashback semanal de 10% sobre perda liquida em mesas Spin na Blaze.",
            "Janela de apuracao: seg 00:00 ate dom 23:59 (horario BR).",
            "Credito em bonus com rollover baixo (1x a 5x) para incentivar retorno.",
            "Teto por jogador e gatilho minimo de perda para proteger margem.",
            "Elegibilidade priorizando UAPI e jogadores com sinal de churn.",
        ],
    )
    card(
        s,
        Inches(6.95),
        Inches(2.0),
        Inches(5.6),
        Inches(3.9),
        "Por que encaixa bem com live casino",
        [
            "Reduz impacto emocional de uma sessao negativa.",
            "Gera motivo claro para voltar na semana seguinte.",
            "Aumenta recorrencia sem canibalizar premio principal da mesa.",
            "Permite personalizacao por segmento (novo, recorrente, VIP).",
        ],
    )

    # Slide 9 - Benchmarks
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Narrativa favoravel com base no mercado",
        "Campanhas de retencao com incentivo recorrente geram reengajamento mensuravel.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(5.9),
        Inches(3.9),
        "Benchmarks relevantes",
        [
            "Cashback no iGaming costuma operar entre 5% e 20% por segmento.",
            "Apostemos (LatAm, Sportradar/VAIX): 85% de reengajamento em jogadores frequentes alvo.",
            "Funstage (Greentube): +199,4% em LTV apos otimizar campanhas CRM.",
            "Kwiff: crescimento de base com melhora simultanea em retencao.",
        ],
    )
    card(
        s,
        Inches(6.9),
        Inches(2.0),
        Inches(5.65),
        Inches(3.9),
        "Leitura para Blaze + Spin",
        [
            "Nao e sobre desconto agressivo, e sobre reduzir churn e elevar valor de ciclo.",
            "Cashback de 10% como ponto de partida preserva margem e aprendizado.",
            "Com dados de coorte, a oferta pode virar alavanca de NGR incremental.",
            "Influencer dealer + cashback cria loop de retorno semanal previsivel.",
        ],
    )
    write_text(
        s,
        Inches(0.7),
        Inches(6.85),
        Inches(12.0),
        Inches(0.35),
        "Fontes: Track360 (glossario iGaming), iGamingToday (Apostemos/VAIX), Xtremepush (cases CRM).",
        size=10,
        color=COL_MUTED,
    )

    # Slide 10 - Simulacao economica
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Simulacao de impacto (exemplo ilustrativo)",
        "Combinacao creator dealer + cashback melhora retorno da coorte.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(3.95),
        Inches(3.8),
        "Cenario base",
        [
            "100 UAPI gerados no ciclo",
            "Deposito medio: R$ 170",
            "Volume medio por UAPI: R$ 1.050",
            "Retencao 30d: 22%",
        ],
    )
    card(
        s,
        Inches(4.9),
        Inches(2.0),
        Inches(3.95),
        Inches(3.8),
        "Com influencer dealer",
        [
            "UAPI +25% via prova social",
            "Deposito medio +12%",
            "Volume por UAPI +18%",
            "Retencao 30d +6 p.p.",
        ],
    )
    card(
        s,
        Inches(9.1),
        Inches(2.0),
        Inches(3.45),
        Inches(3.8),
        "Com cashback 10%",
        [
            "Retorno de churners elegiveis",
            "Mais sessoes por jogador",
            "Maior previsibilidade de receita",
            "Meta: NGR incr. apos promo > 0",
        ],
    )
    write_text(
        s,
        Inches(0.7),
        Inches(6.15),
        Inches(12.0),
        Inches(0.8),
        "Observacao: numeros servem como estrutura de modelagem; calibrar com base real da Blaze (coortes, hold e custo promocional).",
        size=12,
        color=COL_MUTED,
    )

    # Slide 11 - Controles
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Controles para escalar com seguranca",
        "Cashback eficiente depende de regra clara, antifraude e segmentacao.",
    )
    card(
        s,
        Inches(0.7),
        Inches(2.0),
        Inches(3.9),
        Inches(3.8),
        "Regras de produto",
        [
            "Perda liquida minima para ativar cashback",
            "Teto maximo por CPF por semana",
            "Exclusao de contas com perfil de abuso",
        ],
    )
    card(
        s,
        Inches(4.85),
        Inches(2.0),
        Inches(3.9),
        Inches(3.8),
        "Protecao de margem",
        [
            "Acompanhar NGR liquido por segmento",
            "Hold minimo alvo por mesa",
            "A/B test de percentuais (8%, 10%, 12%)",
        ],
    )
    card(
        s,
        Inches(9.0),
        Inches(2.0),
        Inches(3.6),
        Inches(3.8),
        "Compliance e marca",
        [
            "Termos transparentes da campanha",
            "Jogo responsavel em todas as pecas",
            "Governanca conjunta Blaze + Spin",
        ],
    )

    # Slide 12 - Proximos passos
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    add_header(
        s,
        "Proximo passo recomendado",
        "Pilotar, medir e escalar.",
    )
    write_bullets(
        s,
        Inches(0.9),
        Inches(2.15),
        Inches(12.0),
        Inches(3.8),
        [
            "1) Kick-off comercial Blaze + Spin para validar escopo criativo da temporada.",
            "2) Definicao final da mecanica de cashback 10% e regras de elegibilidade.",
            "3) Setup de tracking unico (conteudo -> cadastro -> deposito -> UAPI -> volume).",
            "4) Go-live do piloto e leitura semanal dos KPIs com ajustes taticos.",
            "5) Decisao de escala com base em NGR incremental e custo por UAPI.",
        ],
        size=23,
        spacing=1.08,
    )
    write_text(
        s,
        Inches(0.9),
        Inches(6.65),
        Inches(12.0),
        Inches(0.45),
        "Spin Gaming | Proposta comercial proprietaria para Blaze",
        size=12,
        color=COL_MUTED,
        align=PP_ALIGN.RIGHT,
    )

    out_dir = Path("deliverables")
    out_dir.mkdir(parents=True, exist_ok=True)
    output = out_dir / "Apresentacao_Spin_Blaze_Casting_Cashback.pptx"
    prs.save(output.as_posix())
    print(f"Gerado: {output}")


if __name__ == "__main__":
    build()
