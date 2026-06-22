from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)

BG_DARK = RGBColor(20, 2, 14)
BG_MID = RGBColor(48, 6, 30)
ACCENT = RGBColor(233, 39, 80)
ACCENT_SOFT = RGBColor(145, 20, 58)
WHITE = RGBColor(242, 242, 242)
MUTED = RGBColor(190, 190, 195)


def add_background(slide) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()

    top_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, WIDTH, Inches(0.62))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = BG_MID
    top_band.line.fill.background()

    left_glow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(-0.25), Inches(0.62), Inches(2.4), Inches(6.9)
    )
    left_glow.fill.solid()
    left_glow.fill.fore_color.rgb = ACCENT_SOFT
    left_glow.fill.transparency = 0.6
    left_glow.line.fill.background()

    right_glow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(10.5), Inches(0.62), Inches(2.9), Inches(6.9)
    )
    right_glow.rotation = 180
    right_glow.fill.solid()
    right_glow.fill.fore_color.rgb = ACCENT_SOFT
    right_glow.fill.transparency = 0.65
    right_glow.line.fill.background()

    brand_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.15), Inches(0.1), Inches(4.0), Inches(0.4))
    brand_box.fill.solid()
    brand_box.fill.fore_color.rgb = RGBColor(66, 10, 38)
    brand_box.line.color.rgb = ACCENT
    brand_box.line.width = Pt(1.2)
    brand_tf = brand_box.text_frame
    brand_tf.clear()
    p = brand_tf.paragraphs[0]
    run = p.add_run()
    run.text = "SPIN GAMING  |  BLAZE"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER


def add_title(slide, text: str, subtitle: str | None = None) -> None:
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.7), Inches(1.1))
    tf = title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text.upper()
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    underline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.95), Inches(3.3), Inches(0.05))
    underline.fill.solid()
    underline.fill.fore_color.rgb = ACCENT
    underline.line.fill.background()

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.8), Inches(2.12), Inches(11.4), Inches(0.7))
        sub_tf = sub.text_frame
        sub_tf.clear()
        p2 = sub_tf.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(16)
        r2.font.color.rgb = MUTED
        p2.alignment = PP_ALIGN.LEFT


def add_card(slide, x: float, y: float, w: float, h: float, heading: str, bullets: list[str]) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(38, 9, 26)
    card.line.color.rgb = ACCENT
    card.line.width = Pt(1.2)

    tf = card.text_frame
    tf.clear()
    tf.margin_left = Pt(14)
    tf.margin_right = Pt(14)
    tf.margin_top = Pt(12)

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = heading
    run.font.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = WHITE

    for bullet in bullets:
        bp = tf.add_paragraph()
        bp.text = f"• {bullet}"
        bp.level = 0
        bp.font.size = Pt(13)
        bp.font.color.rgb = MUTED


def add_kpi_table(slide, x: float, y: float, w: float, h: float, rows: list[tuple[str, str]]) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(33, 7, 21)
    panel.line.color.rgb = ACCENT
    panel.line.width = Pt(1.1)

    tf = panel.text_frame
    tf.clear()
    tf.margin_left = Pt(14)
    tf.margin_right = Pt(14)
    tf.margin_top = Pt(10)

    for idx, (metric, value) in enumerate(rows):
        p = tf.add_paragraph() if idx else tf.paragraphs[0]
        p.text = f"{metric}: {value}"
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE if idx == 0 else MUTED


def add_media_placeholder(slide, x: float, y: float, w: float, h: float, label: str) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(27, 8, 18)
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1.0)

    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.size = Pt(14)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Inserir arquivo de mídia nesta área"
    r2.font.size = Pt(11)
    r2.font.color.rgb = MUTED


def build_deck(output_file: Path) -> None:
    prs = Presentation()
    prs.slide_width = WIDTH
    prs.slide_height = HEIGHT
    blank = prs.slide_layouts[6]

    # Slide 1 - Capa
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Proposta Spin Gaming para Blaze", "Influenciador como dealer para gerar UAPI, depósitos e volume de apostas")

    hero = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.6), Inches(11.7), Inches(3.7))
    hero.fill.solid()
    hero.fill.fore_color.rgb = RGBColor(33, 8, 22)
    hero.line.color.rgb = ACCENT
    hero.line.width = Pt(1.2)
    hero_tf = hero.text_frame
    hero_tf.clear()
    p = hero_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "SPIN + BLAZE LIVE CASINO"
    r.font.size = Pt(44)
    r.font.bold = True
    r.font.color.rgb = WHITE

    p2 = hero_tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "PRESENÇA. IDENTIDADE. INFLUÊNCIA."
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT

    # Slide 2 - Contexto
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Por que essa ação é estratégica para a Blaze")

    add_card(
        slide,
        0.8,
        2.3,
        3.9,
        3.8,
        "Marca pronta para escalar presença",
        [
            "A Blaze já tem marca forte e audiência massiva.",
            "O próximo passo é transformar atenção em sessão de jogo.",
            "Live Casino cria conexão em tempo real com o público.",
        ],
    )
    add_card(
        slide,
        4.95,
        2.3,
        3.9,
        3.8,
        "Creator economy acelera conversão",
        [
            "Influenciador reduz atrito de confiança.",
            "Conteúdo social encurta o caminho até o depósito.",
            "Formato gera prova social e compartilhamento orgânico.",
        ],
    )
    add_card(
        slide,
        9.1,
        2.3,
        3.45,
        3.8,
        "Tese central",
        [
            "Treinar um influenciador da marca como dealer cria autenticidade.",
            "Autenticidade aumenta UAPI e recorrência de jogo.",
            "Recorrência eleva depósitos e volume apostado.",
        ],
    )

    # Slide 3 - Case já validado
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Case validado pela Spin: influência que vira audiência")

    add_card(
        slide,
        0.8,
        2.3,
        5.55,
        3.9,
        "Resultados da ação anterior",
        [
            "Quase 1 milhão de visualizações no material principal.",
            "Alto volume de interações e comentários positivos.",
            "Formato com rosto conhecido aumenta tempo de consumo.",
            "Prova de que entretenimento + branding gera tração real.",
        ],
    )
    add_media_placeholder(slide, 6.6, 2.35, 2.75, 1.75, "Vídeo 1 - Massão Dealer")
    add_media_placeholder(slide, 9.55, 2.35, 2.95, 1.75, "Vídeo 2 - Massão Dealer")
    add_media_placeholder(slide, 6.6, 4.35, 5.9, 1.9, "Print da ação de alta visualização")

    # Slide 4 - Proposta de ativação
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Modelo Blaze: Influencer Dealer Program")

    add_card(
        slide,
        0.8,
        2.3,
        12.0,
        3.9,
        "Arquitetura da ação",
        [
            "1) Seleção e treinamento do influenciador como dealer oficial Blaze.",
            "2) Teasers sociais: bastidores + preparação + storytelling de evolução.",
            "3) Sessões ao vivo nas mesas Spin dentro da Blaze com agenda recorrente.",
            "4) Recortes em vídeo para distribuição orgânica e mídia de performance.",
            "5) CTA integrado para entrada nas mesas da campanha.",
        ],
    )

    # Slide 5 - Como gera UAPI
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Por que esse formato aumenta UAPI")

    add_card(
        slide,
        0.8,
        2.3,
        2.85,
        3.8,
        "1. Atenção quente",
        [
            "Audiência já engajada no criador.",
            "Menor custo de atenção.",
        ],
    )
    add_card(
        slide,
        3.95,
        2.3,
        2.85,
        3.8,
        "2. Confiança",
        [
            "Presença humana no produto.",
            "Brand safety e familiaridade.",
        ],
    )
    add_card(
        slide,
        7.1,
        2.3,
        2.85,
        3.8,
        "3. Entrada qualificada",
        [
            "Usuário chega com intenção real.",
            "Maior propensão ao 1º depósito.",
        ],
    )
    add_card(
        slide,
        10.25,
        2.3,
        2.5,
        3.8,
        "4. Recorrência",
        [
            "Volta para novas lives.",
            "Eleva FTD2 e volume.",
        ],
    )

    # Slide 6 - Impacto esperado
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Impacto esperado no funil de receita")

    add_card(
        slide,
        0.8,
        2.3,
        6.1,
        3.9,
        "Premissas de execução",
        [
            "Agenda recorrente de lives com influenciador dealer.",
            "Distribuição com cortes de alta retenção em social.",
            "Promo mecânica conectada ao comportamento da mesa.",
            "Medição por cohort para separar incremento real.",
        ],
    )
    add_kpi_table(
        slide,
        7.15,
        2.3,
        5.4,
        3.9,
        [
            ("Faixas conservadoras para piloto", ""),
            ("UAPI", "+12% a +20%"),
            ("Taxa de 2º depósito", "+8% a +15%"),
            ("Volume de depósitos", "+10% a +18%"),
            ("Volume apostado nas mesas Spin", "+12% a +22%"),
            ("NGR incremental", "+6% a +14%"),
        ],
    )

    # Slide 7 - Cashback
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Programa Cashback 10% para mesas Spin na Blaze")

    add_card(
        slide,
        0.8,
        2.3,
        4.0,
        3.9,
        "Mecânica base",
        [
            "10% de cashback sobre perda líquida nas mesas Spin elegíveis.",
            "Crédito em janela semanal para gerar retorno de sessão.",
            "Comunicação com ranking e sensação de progressão.",
        ],
    )
    add_card(
        slide,
        5.05,
        2.3,
        3.8,
        3.9,
        "Proteção de margem",
        [
            "Teto por jogador no piloto.",
            "Elegibilidade com KYC e histórico saudável.",
            "Regras contra hedge e abuso promocional.",
            "Exclusão de contas sob alerta de risco.",
        ],
    )
    add_card(
        slide,
        9.1,
        2.3,
        3.45,
        3.9,
        "Objetivo de negócio",
        [
            "Aumentar frequência de retorno.",
            "Elevar 2º depósito no mês.",
            "Expandir tempo de mesa e apostas.",
            "Aumentar NGR líquido por cohort.",
        ],
    )

    # Slide 8 - Evidências de mercado
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Por que cashback funciona no mercado de bet")

    add_card(
        slide,
        0.8,
        2.3,
        6.0,
        3.9,
        "Benchmarks e aprendizados",
        [
            "Retenção é o principal alavancador de lucro em modelos transacionais.",
            "Benchmarks iGaming mostram ganho de retenção e depósitos quando loyalty é bem calibrado.",
            "Em mercados maduros, lifecycle marketing reduz dependência de bônus de aquisição.",
            "Cashback funciona melhor quando é direcionado por comportamento (não massivo).",
        ],
    )
    add_card(
        slide,
        7.05,
        2.3,
        5.5,
        3.9,
        "Dados de referência para narrativa",
        [
            "Trueplay (2025): +18% retenção, +10% depósitos, +5,6% NGR em marcas com loyalty.",
            "Optimove Pulse (2025): retenção perto de 70% em mercados globais; CRM é decisivo.",
            "Bain: +5% retenção pode elevar lucro entre 25% e 95% (efeito econômico geral).",
            "Soft2Bet: acompanhar 2º depósito e retenção D30 melhora eficiência promocional.",
        ],
    )

    # Slide 9 - Governança e antifraude
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Governança da campanha: crescimento com controle")

    add_card(
        slide,
        0.8,
        2.3,
        5.8,
        3.9,
        "Controles operacionais",
        [
            "Regras de elegibilidade por CPF, dispositivo e comportamento.",
            "Monitoramento de padrões atípicos de aposta e saque.",
            "Limites dinâmicos de cashback por faixa de risco.",
            "Auditoria semanal entre CRM, risco e operação de jogo.",
        ],
    )
    add_card(
        slide,
        6.9,
        2.3,
        5.65,
        3.9,
        "Medição de incremento real",
        [
            "Grupo holdout sem incentivo para medir lift real.",
            "Leitura por cohort: retenção D7/D30 e 2º depósito.",
            "Acompanhamento NGR/GGR para proteger rentabilidade.",
            "Escalonar somente jornadas com ROI incremental positivo.",
        ],
    )

    # Slide 10 - Plano integrado
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Plano integrado de execução")

    add_card(
        slide,
        0.8,
        2.3,
        2.85,
        3.9,
        "Fase 1",
        [
            "Definir dealer influenciador.",
            "Fechar regras promocionais.",
            "Alinhar tracking completo.",
        ],
    )
    add_card(
        slide,
        3.95,
        2.3,
        2.85,
        3.9,
        "Fase 2",
        [
            "Rodar teaser e bastidores.",
            "Treinamento on-camera.",
            "Campanha de expectativa.",
        ],
    )
    add_card(
        slide,
        7.1,
        2.3,
        2.85,
        3.9,
        "Fase 3",
        [
            "Lives oficiais na Blaze.",
            "Cashback ativo nas mesas.",
            "Criativos de performance.",
        ],
    )
    add_card(
        slide,
        10.25,
        2.3,
        2.5,
        3.9,
        "Fase 4",
        [
            "Análise de cohort.",
            "Ajuste de incentivo.",
            "Escala para novos nomes.",
        ],
    )

    # Slide 11 - KPI de sucesso
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "KPIs para decisão de escala")

    add_card(
        slide,
        0.8,
        2.3,
        6.0,
        3.9,
        "Painel principal",
        [
            "UAPI gerado por origem de tráfego.",
            "Taxa de conversão em 1º e 2º depósito.",
            "Ticket médio e frequência de depósito.",
            "Volume apostado nas mesas Spin da campanha.",
            "NGR incremental e payback promocional.",
        ],
    )
    add_card(
        slide,
        7.05,
        2.3,
        5.5,
        3.9,
        "Critérios de expansão",
        [
            "Escalar se houver lift consistente em retenção e depósitos.",
            "Escalar somente se NGR incremental permanecer saudável.",
            "Abrir novas mesas e novos creators por perfil de audiência.",
            "Transformar ação em calendário fixo da Blaze.",
        ],
    )

    # Slide 12 - Fechamento
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Próximos passos com a Blaze")

    add_card(
        slide,
        0.8,
        2.3,
        12.0,
        3.9,
        "Decisão recomendada",
        [
            "Aprovar piloto Influencer Dealer + Cashback 10% nas mesas Spin.",
            "Validar influenciador, regras de campanha e pacote criativo.",
            "Entrar em operação com medição de incremento desde o dia 1.",
            "Transformar o case em referência de Live Casino no Brasil.",
        ],
    )

    # Slide 13 - Fontes
    slide = prs.slides.add_slide(blank)
    add_background(slide)
    add_title(slide, "Referências de mercado")

    refs = [
        "Optimove iGaming Pulse Report 2025 (US vs Global KPIs).",
        "Optimove - March Madness 2026: retention e qualidade de FTD.",
        "Trueplay reports 2025: impacto de loyalty em retenção, depósitos e NGR.",
        "Soft2Bet: CAC x LTV e eficiência de retenção em iGaming.",
        "Bain & Company / The Loyalty Effect: economia da retenção.",
    ]
    add_card(
        slide,
        0.8,
        2.3,
        12.0,
        3.9,
        "Nota metodológica",
        [
            "Dados de mercado utilizados como benchmark de direção estratégica.",
            "Recomendado validar hipóteses com dados internos Blaze durante o piloto.",
            *refs,
        ],
    )

    prs.save(output_file)


if __name__ == "__main__":
    output_path = Path("Apresentacao_Spin_Blaze_Influencer_Dealer.pptx")
    build_deck(output_path)
    print(f"Deck gerado: {output_path.resolve()}")
